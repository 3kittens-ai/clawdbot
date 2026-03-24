import glob
import os
import sqlite3
from datetime import datetime
from pathlib import Path

import pandas as pd

# 0. 配置
MODEL_ROOT = Path(__file__).resolve().parents[2]
EVAL_PATH = MODEL_ROOT / "outputs" / "backtest_eval.csv"
METRICS_GLOBAL_PATH = MODEL_ROOT / "outputs" / "metrics_global.csv"
METRICS_HORIZON_PATH = MODEL_ROOT / "outputs" / "metrics_horizon.csv"
IMPORTANCE_PATH = MODEL_ROOT / "outputs" / "feature_importance.csv"
DB_PATH = MODEL_ROOT / "data-base" / "sales_filtered.sqlite"
OUTPUT_DIR = MODEL_ROOT / "outputs" / "final"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# 清理旧文件
for old_file in glob.glob(str(OUTPUT_DIR / "db-model-backtest-*.xlsx")):
    try:
        os.remove(old_file)
        print(f"已清理旧回测文件: {os.path.basename(old_file)}")
    except:
        pass

now = datetime.now().strftime('%Y%m%d-%H%M%S')
NEW_EXCEL_PATH = OUTPUT_DIR / f'db-model-backtest-{now}.xlsx'

print(f"解析回测多维数据...")
df = pd.read_csv(EVAL_PATH)
df['variant_key'] = df['variant_key'].astype(str)
metrics_global = pd.read_csv(METRICS_GLOBAL_PATH)
metrics_horizon = pd.read_csv(METRICS_HORIZON_PATH)
feat_imp = pd.read_csv(IMPORTANCE_PATH).head(10) # 仅保留 Top 10

# 2. 构造【说明】页
print("构造【说明】页...")
description_data = [
    ['字段', '值'],
    ['模式', '数据库时间范围回测'],
    ['训练起始月', '2024-04'],
    ['训练截止月', '2025-02'],
    ['评估起始月', '2025-03'],
    ['评估截止月', '2026-02'],
    ['预测月份数', 12],
    ['对比月份', '2025-03 到 2026-02'],
    ['验证策略', 'single_cutoff_holdout'],
    ['步长(月)', 1],
    ['SKU范围', '全部 SKU'],
    ['低预测偏差口径', '实际销量 - 预测销量 (仅统计大于 0 的 SKU)']
]
df_desc = pd.DataFrame(description_data)

# 3. 处理【全量预测结果】页
print("处理【全量预测结果】页...")
conn = sqlite3.connect(DB_PATH)
dim_sku = pd.read_sql("SELECT sku_code as variant_key, sku_name FROM dim_sku", conn)
conn.close()
dim_sku['variant_key'] = dim_sku['variant_key'].astype(str)
df = pd.merge(df, dim_sku, on='variant_key', how='left')

actual_pivot = df.pivot(index='variant_key', columns='year_month', values='monthly_qty').fillna(0)
actual_pivot.columns = [f"{c} 实际" for c in actual_pivot.columns]
forecast_pivot = df.pivot(index='variant_key', columns='year_month', values='forecast_qty').fillna(0)
forecast_pivot.columns = [f"{c} 预测" for c in forecast_pivot.columns]
pivot_df = pd.concat([actual_pivot, forecast_pivot], axis=1)
pivot_df['总实际销量'] = df.groupby('variant_key')['monthly_qty'].sum()
pivot_df['总预测销量'] = df.groupby('variant_key')['forecast_qty'].sum()

merged_df = pd.merge(
    df[['variant_key', 'sku_name', 'product_category', 'family_tags']].drop_duplicates(),
    pivot_df.reset_index(),
    on='variant_key'
)

# 生成 Family 汇总
# 修复 BUG: 显式填充空 family_tags，确保不丢 SKU
merged_df['family_tags'] = merged_df['family_tags'].fillna('未归类')
family_groups = merged_df.groupby('family_tags')
final_rows = []
months = sorted(list(set([c.split(' ')[0] for c in actual_pivot.columns])))
ordered_month_cols = []
for m in months:
    ordered_month_cols.extend([f"{m} 实际", f"{m} 预测"])

family_totals = family_groups['总实际销量'].sum().sort_values(ascending=False)
for family, _ in family_totals.items():
    group = merged_df[merged_df['family_tags'] == family].sort_values('总实际销量', ascending=False)
    summary_row = {
        'variant_key': 'Family TOTAL', 'sku_name': f'【汇总】{family}',
        'product_category': group['product_category'].iloc[0] if not group.empty else '',
        'family_tags': family, '总实际销量': group['总实际销量'].sum(), '总预测销量': group['总预测销量'].sum()
    }
    for col in ordered_month_cols: summary_row[col] = group[col].sum()
    final_rows.append(pd.DataFrame([summary_row]))
    final_rows.append(group)

total_results_df = pd.concat(final_rows, ignore_index=True)
total_results_df = total_results_df.rename(columns={'variant_key': '商品编码', 'sku_name': '商品名称', 'product_category': '产品分类'})
final_cols = ['商品编码', '商品名称', '产品分类', '总实际销量', '总预测销量'] + ordered_month_cols
total_results_df = total_results_df[final_cols]

# 4. 以多 Sheet 方式保存
print(f"正在保存多 Sheet Excel: {NEW_EXCEL_PATH}...")
with pd.ExcelWriter(NEW_EXCEL_PATH, engine='openpyxl') as writer:
    df_desc.to_excel(writer, index=False, header=False, sheet_name='说明')
    metrics_global.to_excel(writer, index=False, sheet_name='整体指标')
    metrics_horizon.to_excel(writer, index=False, sheet_name='分步指标')
    feat_imp.to_excel(writer, index=False, sheet_name='特征 Top10')
    total_results_df.to_excel(writer, index=False, sheet_name='全量预测结果')
    
    from openpyxl.styles import Font, PatternFill, Alignment
    bold_font = Font(bold=True)
    header_fill = PatternFill(start_color='D7E4BC', end_color='D7E4BC', fill_type='solid')
    summary_fill = PatternFill(start_color='FDE9D9', end_color='FDE9D9', fill_type='solid')

    for sheet_name in ['整体指标', '分步指标', '特征 Top10', '全量预测结果']:
        ws = writer.sheets[sheet_name]
        for cell in ws[1]:
            cell.font = bold_font
            cell.fill = header_fill
        
        if sheet_name == '全量预测结果':
            for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                name_val = ws.cell(row=row_idx, column=2).value
                if name_val and '【汇总】' in str(name_val):
                    for cell in row:
                        cell.font = bold_font
                        cell.fill = summary_fill

print("=== 导出完成 (多 Sheet 版) ===")
print(f"FILE_PATH:{NEW_EXCEL_PATH}")
