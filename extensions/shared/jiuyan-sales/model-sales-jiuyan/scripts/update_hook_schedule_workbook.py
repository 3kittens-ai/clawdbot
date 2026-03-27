#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import os
import sqlite3
import tempfile
import zipfile
from collections import OrderedDict
from datetime import datetime, timedelta
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl

MAIN_NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

SOURCE_HEADERS = [
    "商品编码",
    "商品名",
    "产品分类",
    "供应商",
    "实际可用数",
    "库存可售天数",
    "采购在途",
    "自定义实发",
    "7天实发",
    "15天实发",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Refresh hook schedule workbook from SQLite")
    parser.add_argument("--db", required=True)
    parser.add_argument("--source-workbook", required=True)
    parser.add_argument("--plan-workbook", required=True)
    parser.add_argument("--job-id", required=True)
    return parser.parse_args()


def ensure_writable(path: Path) -> None:
    if not path.exists():
        return
    mode = path.stat().st_mode
    if mode & 0o200:
        return
    path.chmod(mode | 0o200)


def load_existing_source_rows(source_workbook: Path) -> tuple[list[str], OrderedDict[str, dict[str, object]]]:
    wb = openpyxl.load_workbook(source_workbook, read_only=True, data_only=False, keep_links=False)
    try:
        ws = wb[wb.sheetnames[0]]
        headers = [cell for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        header_names = [str(value).strip() if value is not None else "" for value in headers]
        rows: OrderedDict[str, dict[str, object]] = OrderedDict()
        for row in ws.iter_rows(min_row=2, values_only=True):
            sku_code = str(row[0] or "").strip()
            if not sku_code:
                continue
            row_map = {
                header_names[index]: row[index] if index < len(row) else None
                for index in range(len(header_names))
            }
            rows[sku_code] = row_map
        return header_names, rows
    finally:
        wb.close()


def load_target_sku_codes(plan_workbook: Path) -> list[str]:
    wb = openpyxl.load_workbook(plan_workbook, read_only=True, data_only=False, keep_links=False)
    try:
        candidate_sheet_names = ["生产计划", "需求预测"]
        sku_codes: list[str] = []
        seen: set[str] = set()
        for sheet_name in candidate_sheet_names:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            start_row = 3 if sheet_name == "生产计划" else 5
            for row in ws.iter_rows(min_row=start_row, max_col=1, values_only=True):
                sku_code = str(row[0] or "").strip()
                if not sku_code or sku_code in seen or not sku_code.isdigit():
                    continue
                seen.add(sku_code)
                sku_codes.append(sku_code)
        return sku_codes
    finally:
        wb.close()


def chunked(sequence: list[str], size: int) -> list[list[str]]:
    return [sequence[index : index + size] for index in range(0, len(sequence), size)]


def query_schedule_metrics(
    db_path: Path,
    sku_codes: list[str],
) -> tuple[str, OrderedDict[str, dict[str, object]]]:
    conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True, timeout=5)
    conn.row_factory = sqlite3.Row
    try:
        latest_date_row = conn.execute("SELECT MAX(sale_date) AS latest_date FROM sales").fetchone()
        latest_date = str(latest_date_row["latest_date"] or "").strip()
        if not latest_date:
            raise RuntimeError("sales 数据库为空，无法刷新排单表")

        month_start = f"{latest_date[:7]}-01"
        latest_date_obj = datetime.strptime(latest_date, "%Y-%m-%d").date()
        recent_15_start = (latest_date_obj - timedelta(days=14)).strftime("%Y-%m-%d")
        lookback_start = min(month_start, recent_15_start)
        metrics: OrderedDict[str, dict[str, object]] = OrderedDict()
        cleaned_sku_codes = [sku_code for sku_code in sku_codes if sku_code]
        for sku_chunk in chunked(cleaned_sku_codes, 200):
            placeholders = ",".join("?" for _ in sku_chunk)
            rows = conn.execute(
                f"""
                SELECT
                  s.barcode AS sku_code,
                  COALESCE(d.sku_name, MAX(s.product_name), '') AS sku_name,
                  COALESCE(d.product_category, MAX(s.product_category), '') AS product_category,
                  SUM(CASE WHEN s.sale_date >= ? THEN s.sales_volume ELSE 0 END) AS current_month_qty,
                  SUM(CASE WHEN s.sale_date >= date(?, '-6 day') THEN s.sales_volume ELSE 0 END) AS qty_7d,
                  SUM(CASE WHEN s.sale_date >= date(?, '-14 day') THEN s.sales_volume ELSE 0 END) AS qty_15d
                FROM sales s
                LEFT JOIN dim_sku d ON d.sku_code = s.barcode
                WHERE s.barcode IN ({placeholders})
                  AND s.sale_date >= ?
                GROUP BY s.barcode
                ORDER BY s.barcode
                """,
                (month_start, latest_date, latest_date, *sku_chunk, lookback_start),
            ).fetchall()
            for row in rows:
                sku_code = str(row["sku_code"] or "").strip()
                if not sku_code:
                    continue
                metrics[sku_code] = {
                    "商品编码": sku_code,
                    "商品名": row["sku_name"] or "",
                    "产品分类": row["product_category"] or "",
                    "自定义实发": int(row["current_month_qty"] or 0),
                    "7天实发": int(row["qty_7d"] or 0),
                    "15天实发": int(row["qty_15d"] or 0),
                }
        return latest_date, metrics
    finally:
        conn.close()


def build_source_rows(
    target_sku_codes: list[str],
    existing_rows: OrderedDict[str, dict[str, object]],
    metrics_by_sku: OrderedDict[str, dict[str, object]],
) -> list[dict[str, object]]:
    merged_rows: list[dict[str, object]] = []
    for sku_code in target_sku_codes:
        existing = existing_rows.get(sku_code, {})
        metrics = metrics_by_sku.get(sku_code, {})
        merged_rows.append(
            {
                "商品编码": sku_code,
                "商品名": metrics.get("商品名") or existing.get("商品名") or "",
                "产品分类": metrics.get("产品分类") or existing.get("产品分类") or "",
                "供应商": existing.get("供应商") or "",
                "实际可用数": existing.get("实际可用数") or 0,
                "库存可售天数": existing.get("库存可售天数") or "",
                "采购在途": existing.get("采购在途") or 0,
                "自定义实发": metrics.get("自定义实发", 0),
                "7天实发": metrics.get("7天实发", 0),
                "15天实发": metrics.get("15天实发", 0),
            }
        )
    return merged_rows


def write_source_workbook(source_workbook: Path, rows: list[dict[str, object]]) -> None:
    ensure_writable(source_workbook)
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "Sheet1"
    ws.append(SOURCE_HEADERS)
    for row in rows:
        ws.append([row.get(header, "") for header in SOURCE_HEADERS])
    workbook.save(source_workbook)
    workbook.close()


def column_letter(index: int) -> str:
    result = ""
    current = index
    while current > 0:
        current, remainder = divmod(current - 1, 26)
        result = chr(65 + remainder) + result
    return result


def build_external_link_xml(rows: list[dict[str, object]]) -> bytes:
    ET.register_namespace("", MAIN_NAMESPACE)
    ET.register_namespace("r", REL_NAMESPACE)
    external_link = ET.Element(f"{{{MAIN_NAMESPACE}}}externalLink")
    external_book = ET.SubElement(
        external_link,
        f"{{{MAIN_NAMESPACE}}}externalBook",
        {f"{{{REL_NAMESPACE}}}id": "rId1"},
    )
    sheet_names = ET.SubElement(external_book, f"{{{MAIN_NAMESPACE}}}sheetNames")
    ET.SubElement(sheet_names, f"{{{MAIN_NAMESPACE}}}sheetName", {"val": "Sheet1"})
    sheet_data_set = ET.SubElement(external_book, f"{{{MAIN_NAMESPACE}}}sheetDataSet")
    sheet_data = ET.SubElement(sheet_data_set, f"{{{MAIN_NAMESPACE}}}sheetData", {"sheetId": "0"})

    header_row = ET.SubElement(sheet_data, f"{{{MAIN_NAMESPACE}}}row", {"r": "1"})
    for column_index, header in enumerate(SOURCE_HEADERS, start=1):
        cell = ET.SubElement(
            header_row,
            f"{{{MAIN_NAMESPACE}}}cell",
            {"r": f"{column_letter(column_index)}1", "t": "str"},
        )
        ET.SubElement(cell, f"{{{MAIN_NAMESPACE}}}v").text = str(header)

    for row_index, row in enumerate(rows, start=2):
        xml_row = ET.SubElement(sheet_data, f"{{{MAIN_NAMESPACE}}}row", {"r": str(row_index)})
        for column_index, header in enumerate(SOURCE_HEADERS, start=1):
            value = row.get(header, "")
            if value in (None, ""):
                continue
            attrs = {"r": f"{column_letter(column_index)}{row_index}"}
            if isinstance(value, str):
                attrs["t"] = "str"
            cell = ET.SubElement(xml_row, f"{{{MAIN_NAMESPACE}}}cell", attrs)
            ET.SubElement(cell, f"{{{MAIN_NAMESPACE}}}v").text = str(value)

    return ET.tostring(external_link, encoding="utf-8", xml_declaration=True)


def rewrite_zip_entry(archive_path: Path, replacements: dict[str, bytes]) -> None:
    ensure_writable(archive_path)
    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as temp_file:
        temp_path = Path(temp_file.name)

    try:
        with zipfile.ZipFile(archive_path, "r") as source_zip, zipfile.ZipFile(
            temp_path,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as target_zip:
            for info in source_zip.infolist():
                data = replacements.get(info.filename)
                if data is None:
                    data = source_zip.read(info.filename)
                target_zip.writestr(info, data)
        os.replace(temp_path, archive_path)
    finally:
        temp_path.unlink(missing_ok=True)


def update_plan_workbook(plan_workbook: Path, rows: list[dict[str, object]]) -> None:
    rewrite_zip_entry(
        plan_workbook,
        {"xl/externalLinks/externalLink9.xml": build_external_link_xml(rows)},
    )


def main() -> None:
    args = parse_args()
    db_path = Path(args.db)
    source_workbook = Path(args.source_workbook)
    plan_workbook = Path(args.plan_workbook)

    target_sku_codes = load_target_sku_codes(plan_workbook)
    _, existing_rows = load_existing_source_rows(source_workbook)
    latest_date, metrics_by_sku = query_schedule_metrics(db_path, target_sku_codes)
    rows = build_source_rows(target_sku_codes, existing_rows, metrics_by_sku)
    write_source_workbook(source_workbook, rows)
    update_plan_workbook(plan_workbook, rows)

    print(
        json.dumps(
            {
                "file_path": str(plan_workbook),
                "file_name": plan_workbook.name,
                "source_file_path": str(source_workbook),
                "source_file_name": source_workbook.name,
                "latest_date": latest_date,
                "updated_sku_count": len(metrics_by_sku),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
