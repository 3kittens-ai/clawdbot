#!/usr/bin/env python3
from __future__ import annotations

import argparse
import ast
import json
import math
import re
import sqlite3
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

FINAL_OUTPUT_KEEP_COUNT = 50
BUILTIN_DEFAULT_FORMULA = "builtin_default_5m"


@dataclass
class LoadedSalesData:
    monthly_df: pd.DataFrame
    daily_df: pd.DataFrame
    latest_sale_date: str
    latest_sale_month: str
    latest_complete_month: str
    inventory_metrics_updated_at: str | None


def cleanup_final_outputs(output_dir: Path, keep_count: int = FINAL_OUTPUT_KEEP_COUNT) -> None:
    files = [path for path in output_dir.iterdir() if path.is_file()]
    if len(files) <= keep_count:
        return
    files.sort(key=lambda path: path.stat().st_mtime, reverse=True)
    for old_file in files[keep_count:]:
        old_file.unlink(missing_ok=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate formula-based sales forecast Excel")
    parser.add_argument("--db", required=True)
    parser.add_argument("--formula", required=True)
    parser.add_argument("--formula-text", required=True)
    parser.add_argument("--months", type=int, required=True)
    parser.add_argument("--top-k", type=int, required=True)
    parser.add_argument("--target-turnover-days", type=int, default=45)
    parser.add_argument(
        "--top-basis",
        choices=["latest_day", "latest_complete_month", "trailing_12_months", "all_time"],
        default="latest_day",
    )
    parser.add_argument("--sku-list-excel")
    parser.add_argument("--product-category")
    parser.add_argument("--job-id", required=True)
    return parser.parse_args()


def month_shift(year_month: str, offset: int) -> str:
    year, month = [int(part) for part in year_month.split("-")]
    total = (year * 12 + (month - 1)) + offset
    next_year = total // 12
    next_month = total % 12 + 1
    return f"{next_year:04d}-{next_month:02d}"


def month_end(year_month: str) -> str:
    next_month = month_shift(year_month, 1)
    next_month_start = datetime.strptime(f"{next_month}-01", "%Y-%m-%d").date()
    return (next_month_start - timedelta(days=1)).strftime("%Y-%m-%d")


def resolve_latest_complete_month(latest_sale_date: str) -> str:
    latest_sale_day = datetime.strptime(latest_sale_date, "%Y-%m-%d").date()
    latest_month = latest_sale_date[:7]
    if latest_sale_date == month_end(latest_month):
        return latest_month
    return month_shift(latest_month, -1)


def build_safe_env(history: Sequence[float]) -> dict[str, float]:
    latest = list(history)[-12:]
    padded = [0.0] * max(0, 12 - len(latest)) + latest

    def avg(count: int) -> float:
        bucket = padded[-count:]
        return float(sum(bucket) / count) if count else 0.0

    def total(count: int) -> float:
        return float(sum(padded[-count:])) if count else 0.0

    m1 = float(padded[-1])
    m2 = float(padded[-2])
    return {
        **{f"m{i}": float(padded[-i]) for i in range(1, 13)},
        "avg3": avg(3),
        "avg6": avg(6),
        "avg12": avg(12),
        "sum3": total(3),
        "sum6": total(6),
        "sum12": total(12),
        "trend": float(m1 - m2),
        "growth": float((m1 - m2) / m2) if m2 else 0.0,
    }


ALLOWED_FUNCS = {
    "min": min,
    "max": max,
    "round": round,
}


def eval_formula(expr: str, env: dict[str, float]) -> float:
    node = ast.parse(expr, mode="eval")

    def _eval(current: ast.AST) -> float:
        if isinstance(current, ast.Expression):
            return _eval(current.body)
        if isinstance(current, ast.Constant) and isinstance(current.value, (int, float)):
            return float(current.value)
        if isinstance(current, ast.Name):
            if current.id not in env:
                raise ValueError(f"unsupported variable: {current.id}")
            return float(env[current.id])
        if isinstance(current, ast.BinOp):
            left = _eval(current.left)
            right = _eval(current.right)
            if isinstance(current.op, ast.Add):
                return left + right
            if isinstance(current.op, ast.Sub):
                return left - right
            if isinstance(current.op, ast.Mult):
                return left * right
            if isinstance(current.op, ast.Div):
                return left / right if right else 0.0
            if isinstance(current.op, ast.Pow):
                return left**right
            raise ValueError("unsupported operator")
        if isinstance(current, ast.UnaryOp):
            operand = _eval(current.operand)
            if isinstance(current.op, ast.USub):
                return -operand
            if isinstance(current.op, ast.UAdd):
                return operand
            raise ValueError("unsupported unary operator")
        if isinstance(current, ast.Call) and isinstance(current.func, ast.Name):
            func = ALLOWED_FUNCS.get(current.func.id)
            if func is None:
                raise ValueError(f"unsupported function: {current.func.id}")
            args = [_eval(arg) for arg in current.args]
            return float(func(*args))
        raise ValueError("unsupported expression")

    value = _eval(node)
    if not math.isfinite(value):
        raise ValueError("formula result is not finite")
    return max(value, 0.0)


def build_top_sku_filter(
    top_basis: str,
    latest_sale_date: str,
    latest_complete_month: str,
) -> tuple[str, list[str]]:
    if top_basis == "all_time":
        return "", []
    if top_basis == "trailing_12_months":
        start_month = month_shift(latest_complete_month, -11)
        return "WHERE sale_date >= ?", [f"{start_month}-01"]
    if top_basis == "latest_complete_month":
        month_start = f"{latest_complete_month}-01"
        month_stop = f"{month_shift(latest_complete_month, 1)}-01"
        return "WHERE sale_date >= ? AND sale_date < ?", [month_start, month_stop]
    return "WHERE sale_date = ?", [latest_sale_date]


def load_top_history(
    db_path: str,
    top_k: int,
    top_basis: str = "latest_day",
    product_category: str | None = None,
    sku_filter_codes: Sequence[str] | None = None,
) -> LoadedSalesData:
    conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True, timeout=5)
    conn.execute("PRAGMA query_only = ON")
    dim_sku_columns = {
        str(row[1])
        for row in conn.execute("PRAGMA table_info(dim_sku)").fetchall()
    }
    latest_inventory_select = (
        "d.latest_inventory AS latest_inventory"
        if "latest_inventory" in dim_sku_columns
        else "NULL AS latest_inventory"
    )
    latest_in_transit_select = (
        "d.latest_in_transit AS latest_in_transit"
        if "latest_in_transit" in dim_sku_columns
        else "NULL AS latest_in_transit"
    )
    latest_inventory_group_by = "d.latest_inventory" if "latest_inventory" in dim_sku_columns else "NULL"
    latest_in_transit_group_by = "d.latest_in_transit" if "latest_in_transit" in dim_sku_columns else "NULL"
    inventory_metrics_updated_at: str | None = None
    if {"latest_inventory", "latest_in_transit"}.intersection(dim_sku_columns):
        inventory_metrics_updated_at = pd.read_sql_query(
            """
            SELECT MAX(updated_at) AS inventory_metrics_updated_at
            FROM dim_sku
            WHERE latest_inventory IS NOT NULL OR latest_in_transit IS NOT NULL
            """,
            conn,
        ).iloc[0]["inventory_metrics_updated_at"]
        if pd.notna(inventory_metrics_updated_at):
            inventory_metrics_updated_at = str(inventory_metrics_updated_at)
        else:
            inventory_metrics_updated_at = None
    latest_sale_date = pd.read_sql_query(
        "SELECT MAX(sale_date) AS latest_sale_date FROM sales",
        conn,
    ).iloc[0]["latest_sale_date"]
    latest_sale_date = str(latest_sale_date)
    latest_sale_month = latest_sale_date[:7]
    latest_complete_month = resolve_latest_complete_month(latest_sale_date)
    history_start = month_shift(latest_sale_month, -13)
    history_start_date = f"{history_start}-01"
    current_month_start = f"{latest_sale_month}-01"
    top_basis_where_sql, top_basis_params = build_top_sku_filter(
        top_basis,
        latest_sale_date,
        latest_complete_month,
    )
    recent_daily_start = (
        datetime.strptime(latest_sale_date, "%Y-%m-%d").date() - timedelta(days=14)
    ).strftime("%Y-%m-%d")
    daily_query_start = min(current_month_start, recent_daily_start)

    if sku_filter_codes:
        selected_skus_rows = ",\n          ".join("(?, ?)" for _ in sku_filter_codes)
        selected_skus_cte = f"""
        WITH selected_skus(sku_code, sort_order) AS (
          VALUES
          {selected_skus_rows}
        )
        """
        selected_sku_params: list[object] = []
        for sort_order, sku_code in enumerate(sku_filter_codes):
            selected_sku_params.extend([sku_code, sort_order])
        monthly_sql = (
            selected_skus_cte
            + """
        SELECT
          s.sku_code,
          d.sku_name,
          d.variant_key AS sku_variant_key,
          """
            + latest_inventory_select
            + """
          ,
          """
            + latest_in_transit_select
            + """
          ,
          substr(r.sale_date, 1, 7) AS year_month,
          SUM(r.sales_volume) AS monthly_qty,
          s.sort_order AS total_qty
        FROM selected_skus s
        LEFT JOIN dim_sku d ON d.sku_code = s.sku_code
        LEFT JOIN sales r INDEXED BY idx_sales_barcode_date
          ON r.barcode = s.sku_code
         AND r.sale_date >= ?
        """
            + (" AND r.product_category = ?" if product_category else "")
            + """
        GROUP BY s.sku_code, d.sku_name, d.variant_key, """
            + latest_inventory_group_by
            + """, """
            + latest_in_transit_group_by
            + """, substr(r.sale_date, 1, 7), s.sort_order
        ORDER BY s.sort_order ASC, year_month
        """
        )
        monthly_params = selected_sku_params + [history_start_date]
        if product_category:
            monthly_params.append(product_category)
        daily_sql = (
            selected_skus_cte
            + """
        SELECT
          r.barcode AS sku_code,
          r.sale_date,
          SUM(r.sales_volume) AS daily_qty
        FROM selected_skus s
        INNER JOIN sales r INDEXED BY idx_sales_barcode_date
          ON r.barcode = s.sku_code
         AND r.sale_date >= ?
        """
            + (" AND r.product_category = ?" if product_category else "")
            + """
        GROUP BY r.barcode, r.sale_date, s.sort_order
        ORDER BY s.sort_order ASC, r.sale_date
        """
        )
        daily_params = selected_sku_params + [daily_query_start]
        if product_category:
            daily_params.append(product_category)
    elif product_category:
        top_skus_cte = (
            """
        WITH top_skus AS (
          SELECT barcode AS sku_code, SUM(sales_volume) AS total_qty
          FROM sales INDEXED BY idx_sales_date
        """
            + top_basis_where_sql
            + "\n"
            + (" AND product_category = ?" if top_basis_where_sql else "WHERE product_category = ?")
            + """
          GROUP BY barcode
          ORDER BY total_qty DESC
          LIMIT ?
        )
        """
        )
        monthly_sql = (
            top_skus_cte
            + """
        SELECT
          t.sku_code,
          d.sku_name,
          d.variant_key AS sku_variant_key,
          """
            + latest_inventory_select
            + """
          ,
          """
            + latest_in_transit_select
            + """
          ,
          substr(r.sale_date, 1, 7) AS year_month,
          SUM(r.sales_volume) AS monthly_qty,
          t.total_qty
        FROM top_skus t
        LEFT JOIN dim_sku d ON d.sku_code = t.sku_code
        LEFT JOIN sales r INDEXED BY idx_sales_barcode_date
          ON r.barcode = t.sku_code
         AND r.sale_date >= ?
         AND r.product_category = ?
        GROUP BY t.sku_code, d.sku_name, d.variant_key, """
            + latest_inventory_group_by
            + """, """
            + latest_in_transit_group_by
            + """, substr(r.sale_date, 1, 7), t.total_qty
        ORDER BY t.total_qty DESC, t.sku_code, year_month
        """
        )
        monthly_params = top_basis_params + [product_category, top_k, history_start_date, product_category]
        daily_sql = (
            top_skus_cte
            + """
        SELECT
          r.barcode AS sku_code,
          r.sale_date,
          SUM(r.sales_volume) AS daily_qty
        FROM top_skus t
        INNER JOIN sales r INDEXED BY idx_sales_barcode_date
          ON r.barcode = t.sku_code
         AND r.sale_date >= ?
        GROUP BY r.barcode, r.sale_date
        ORDER BY r.barcode, r.sale_date
        """
        )
        daily_params = top_basis_params + [product_category, top_k, daily_query_start]
    else:
        top_skus_cte = (
            """
        WITH top_skus AS (
          SELECT barcode AS sku_code, SUM(sales_volume) AS total_qty
          FROM sales INDEXED BY idx_sales_date
        """
            + top_basis_where_sql
            + """
          GROUP BY barcode
          ORDER BY total_qty DESC
          LIMIT ?
        )
        """
        )
        monthly_sql = (
            top_skus_cte
            + """
        SELECT
          t.sku_code,
          d.sku_name,
          d.variant_key AS sku_variant_key,
          """
            + latest_inventory_select
            + """
          ,
          """
            + latest_in_transit_select
            + """
          ,
          substr(r.sale_date, 1, 7) AS year_month,
          SUM(r.sales_volume) AS monthly_qty,
          t.total_qty
        FROM top_skus t
        LEFT JOIN dim_sku d ON d.sku_code = t.sku_code
        LEFT JOIN sales r INDEXED BY idx_sales_barcode_date
          ON r.barcode = t.sku_code
         AND r.sale_date >= ?
        GROUP BY t.sku_code, d.sku_name, d.variant_key, """
            + latest_inventory_group_by
            + """, """
            + latest_in_transit_group_by
            + """, substr(r.sale_date, 1, 7), t.total_qty
        ORDER BY t.total_qty DESC, t.sku_code, year_month
        """
        )
        monthly_params = top_basis_params + [top_k, history_start_date]
        daily_sql = (
            top_skus_cte
            + """
        SELECT
          r.barcode AS sku_code,
          r.sale_date,
          SUM(r.sales_volume) AS daily_qty
        FROM top_skus t
        INNER JOIN sales r INDEXED BY idx_sales_barcode_date
          ON r.barcode = t.sku_code
         AND r.sale_date >= ?
        GROUP BY r.barcode, r.sale_date
        ORDER BY r.barcode, r.sale_date
        """
        )
        daily_params = top_basis_params + [top_k, daily_query_start]

    monthly_df = pd.read_sql_query(
        monthly_sql,
        conn,
        params=monthly_params,
    )
    daily_df = pd.read_sql_query(
        daily_sql,
        conn,
        params=daily_params,
    )
    conn.close()
    return LoadedSalesData(
        monthly_df=monthly_df,
        daily_df=daily_df,
        latest_sale_date=latest_sale_date,
        latest_sale_month=latest_sale_month,
        latest_complete_month=latest_complete_month,
        inventory_metrics_updated_at=inventory_metrics_updated_at,
    )


def extract_product_root(raw_variant_key: object) -> str:
    value = str(raw_variant_key or "").strip()
    if not value:
        return "未分类"
    parts = value.split("|")
    if len(parts) <= 4:
        return value
    root = "|".join(parts[:-4]).strip("|")
    return root or value


def append_product_root_summary_rows(
    ordered: pd.DataFrame,
    sum_columns: list[str],
    month_series_columns: list[str],
    future_months: list[str],
    total_column: str,
) -> pd.DataFrame:
    if ordered.empty:
        return ordered

    root_column = "商品标签分类"
    code_column = "商品编码"
    name_column = "商品名称"
    ordered = ordered.copy()
    ordered[root_column] = ordered[root_column].fillna("未分类").replace("", "未分类")

    root_totals = (
        ordered.groupby(root_column, as_index=False)[total_column].sum().sort_values(
            by=[total_column, root_column],
            ascending=[False, True],
        )
    )
    sections: list[pd.DataFrame] = []
    for product_root in root_totals[root_column]:
        group = ordered[ordered[root_column] == product_root].copy()
        group = group.sort_values(by=[total_column, code_column], ascending=[False, True])
        summary_row: dict[str, object] = {
            root_column: product_root,
            code_column: "汇总",
            name_column: f"{product_root} 汇总",
        }
        for sum_column in sum_columns:
            summary_row[sum_column] = float(group[sum_column].sum())
        for future_month in future_months:
            previous_month = month_shift(future_month, -1)
            previous_qty = float(summary_row.get(previous_month, 0.0) or 0.0)
            current_qty = float(summary_row.get(future_month, 0.0) or 0.0)
            summary_row[f"{future_month}环比增幅(%)"] = format_pct_display(
                compute_mom_growth_pct(current_qty, previous_qty)
            )
        growth_row: dict[str, object] = {
            root_column: product_root,
            code_column: "环比",
            name_column: f"{product_root} 汇总环比",
        }
        yoy_row: dict[str, object] = {
            root_column: product_root,
            code_column: "同比",
            name_column: f"{product_root} 汇总同比",
        }
        previous_month_qty: float | None = None
        for month_column in month_series_columns:
            current_qty = float(summary_row.get(month_column, 0.0) or 0.0)
            growth_row[month_column] = (
                format_pct_display(compute_mom_growth_pct(current_qty, previous_month_qty))
                if previous_month_qty is not None
                else None
            )
            last_year_month = month_shift(month_column, -12)
            last_year_qty = summary_row.get(last_year_month)
            yoy_row[month_column] = (
                format_pct_display(compute_mom_growth_pct(current_qty, float(last_year_qty or 0.0)))
                if last_year_qty is not None
                else None
            )
            previous_month_qty = current_qty
        for future_month in future_months:
            growth_row[f"{future_month}环比增幅(%)"] = summary_row.get(f"{future_month}环比增幅(%)")
        sections.append(group)
        sections.append(pd.DataFrame([summary_row], columns=ordered.columns))
        sections.append(pd.DataFrame([growth_row], columns=ordered.columns))
        sections.append(pd.DataFrame([yoy_row], columns=ordered.columns))

    return pd.concat(sections, ignore_index=True)


def build_month_sequence(start_month: str, end_month: str) -> list[str]:
    if start_month > end_month:
        return []
    months: list[str] = []
    current = start_month
    while current <= end_month:
        months.append(current)
        current = month_shift(current, 1)
    return months


def build_actual_summary_column_names(latest_complete_month: str) -> tuple[str, str, str]:
    current_year = latest_complete_month[:4]
    previous_year = str(int(current_year) - 1)
    two_years_ago = str(int(current_year) - 2)
    return (
        f"{two_years_ago}年实际销量汇总",
        f"{previous_year}年实际销量汇总",
        f"{current_year}年销量汇总",
    )


def build_current_month_mean_forecast_column_name(latest_sale_month: str, latest_complete_month: str) -> str | None:
    if latest_sale_month == latest_complete_month:
        return None
    return f"{latest_sale_month}均值预测销量"


def compute_current_month_mean_forecast_qty(
    sku_daily: list[tuple[str, float]],
    latest_sale_month: str,
) -> int:
    current_month_days = [
        sale_date
        for sale_date, qty in sku_daily
        if sale_date.startswith(latest_sale_month) and float(qty or 0.0) > 0.0
    ]
    sales_day_count = len(current_month_days)
    if sales_day_count <= 0:
        return 0

    current_month_actual = sum(
        float(qty or 0.0)
        for sale_date, qty in sku_daily
        if sale_date.startswith(latest_sale_month)
    )
    month_day_count = int(month_end(latest_sale_month)[-2:])
    return ceil_non_negative(current_month_actual / sales_day_count * month_day_count)


def compute_actual_year_summaries(
    month_map: dict[str, float],
    actual_history_months: list[str],
    latest_complete_month: str,
    latest_sale_month: str,
    current_month_actual: float,
) -> tuple[float, float, float]:
    current_year = latest_complete_month[:4]
    previous_year = str(int(current_year) - 1)
    two_years_ago = str(int(current_year) - 2)
    two_years_ago_total = round(
        sum(float(month_map.get(month, 0.0)) for month in actual_history_months if month.startswith(two_years_ago)),
        4,
    )
    previous_year_total = round(
        sum(float(month_map.get(month, 0.0)) for month in actual_history_months if month.startswith(previous_year)),
        4,
    )
    current_year_total = round(
        sum(float(month_map.get(month, 0.0)) for month in actual_history_months if month.startswith(current_year)),
        4,
    )
    if latest_sale_month.startswith(current_year) and latest_sale_month != latest_complete_month:
        current_year_total = round(current_year_total + current_month_actual, 4)
    return two_years_ago_total, previous_year_total, current_year_total


def safe_ratio(numerator: float, denominator: float, fallback: float = 1.0) -> float:
    if denominator <= 0:
        return fallback
    return numerator / denominator


def ceil_non_negative(value: float) -> int:
    return max(0, int(math.ceil(value)))


def ceil_to_batch(value: float, batch_size: int) -> int:
    if batch_size <= 0:
        return ceil_non_negative(value)
    return max(0, int(math.ceil(max(value, 0.0) / batch_size) * batch_size))


def normalize_sku_header(value: object) -> str:
    return str(value or "").strip().replace("\n", "").replace(" ", "").replace("\u3000", "").lower()


def is_valid_sku_code(value: object) -> bool:
    sku_code = str(value or "").strip()
    return bool(re.fullmatch(r"69\d{11}", sku_code))


def load_sku_codes_from_excel(excel_path: str) -> list[str]:
    try:
        workbook = pd.ExcelFile(excel_path)
    except Exception as exc:  # pragma: no cover - surfaced to caller as job failure
        raise ValueError(f"读取 SKU Excel 失败：{exc}") from exc

    sku_codes: list[str] = []
    seen: set[str] = set()
    normalized_aliases = {
        normalize_sku_header(alias)
        for alias in ("商品编码", "sku_code", "barcode", "SKU", "SKU编码", "69码", "货号", "条码", "单品编码")
    }

    for sheet_name in workbook.sheet_names:
        try:
            sheet_df = pd.read_excel(workbook, sheet_name=sheet_name, header=None)
        except Exception:
            continue
        if sheet_df.empty:
            continue

        header_row_index: int | None = None
        code_column_index: int | None = None
        scan_row_count = min(len(sheet_df.index), 8)
        for row_index in range(scan_row_count):
            row_values = sheet_df.iloc[row_index].tolist()
            for column_index, raw_header in enumerate(row_values):
                if normalize_sku_header(raw_header) in normalized_aliases:
                    header_row_index = row_index
                    code_column_index = column_index
                    break
            if code_column_index is not None:
                break

        if code_column_index is None or header_row_index is None:
            continue

        for raw_value in sheet_df.iloc[header_row_index + 1 :, code_column_index].tolist():
            if pd.isna(raw_value):
                continue
            sku_code = str(raw_value).strip()
            if normalize_sku_header(sku_code) in normalized_aliases:
                continue
            if not is_valid_sku_code(sku_code):
                continue
            if not sku_code or sku_code in seen:
                continue
            seen.add(sku_code)
            sku_codes.append(sku_code)

        if sku_codes:
            return sku_codes

    if not sku_codes:
        raise ValueError("SKU Excel 中未找到有效的 SKU 编码（仅支持 69 开头的 13 位数字）")
    return sku_codes


def format_pct_display(value: float | None) -> str | None:
    if value is None:
        return None
    return f"{int(round(value))}%"


def compute_mom_growth_pct(current_qty: float, previous_qty: float) -> float | None:
    if previous_qty <= 0:
        return None
    return round((current_qty - previous_qty) / previous_qty * 100.0, 2)


def compute_recent_12m_max_qty(month_map: dict[str, float], anchor_month: str) -> float:
    recent_months = [month_shift(anchor_month, -offset) for offset in range(1, 13)]
    return round(max((float(month_map.get(month, 0.0)) for month in recent_months), default=0.0), 4)


def compute_recent_12m_avg_qty(month_map: dict[str, float], anchor_month: str) -> float:
    recent_months = [month_shift(anchor_month, -offset) for offset in range(1, 13)]
    values = [float(month_map.get(month, 0.0)) for month in recent_months]
    return round(sum(values) / len(values), 4) if values else 0.0


def compute_short_term_reference(recent_7_day_qty: float, recent_15_day_qty: float, latest_sale_month: str) -> float:
    month_day_count = int(month_end(latest_sale_month)[-2:])
    return round(max(recent_7_day_qty / 7.0, recent_15_day_qty / 15.0) * month_day_count, 4)


def compute_last_month_yoy_mom_reference(month_map: dict[str, float], latest_sale_month: str, latest_complete_month: str) -> float:
    target_month_last_year = month_shift(latest_sale_month, -12)
    latest_complete_month_last_year = month_shift(latest_complete_month, -12)
    return round(
        month_map.get(latest_complete_month, 0.0)
        * safe_ratio(
            month_map.get(target_month_last_year, 0.0),
            month_map.get(latest_complete_month_last_year, 0.0),
            fallback=1.0,
        ),
        4,
    )


def build_default_month1_forecast(
    month_map: dict[str, float],
    current_month_actual: float,
    recent_7_day_qty: float,
    recent_15_day_qty: float,
    latest_sale_date: str,
    latest_sale_month: str,
    latest_complete_month: str,
) -> float:
    latest_sale_day = datetime.strptime(latest_sale_date, "%Y-%m-%d").date()
    month_day_count = int(month_end(latest_sale_month)[-2:])
    elapsed_days = latest_sale_day.day
    target_month_last_year = month_shift(latest_sale_month, -12)
    latest_complete_month_last_year = month_shift(latest_complete_month, -12)

    historical_ratio_value = month_map.get(latest_complete_month, 0.0) * safe_ratio(
        month_map.get(target_month_last_year, 0.0),
        month_map.get(latest_complete_month_last_year, 0.0),
        fallback=1.0,
    )
    current_month_run_rate_value = (
        current_month_actual / elapsed_days * month_day_count if elapsed_days else 0.0
    )
    short_term_run_rate_value = (
        max(recent_7_day_qty / 7.0, recent_15_day_qty / 15.0) * month_day_count
    )
    return float(
        ceil_non_negative(
        max(
            historical_ratio_value,
            current_month_run_rate_value,
            short_term_run_rate_value,
        )
    ))


def generate_builtin_default_forecast(data: LoadedSalesData, months: int) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    all_future_months = [month_shift(data.latest_sale_month, step) for step in range(months)]
    actual_history_months = build_month_sequence("2024-01", data.latest_complete_month)
    current_month_mean_forecast_column = build_current_month_mean_forecast_column_name(
        data.latest_sale_month, data.latest_complete_month
    )
    two_years_ago_total_column, previous_year_total_column, current_year_total_column = build_actual_summary_column_names(
        data.latest_complete_month
    )

    daily_records: dict[str, list[tuple[str, float]]] = {}
    for _, row in data.daily_df.iterrows():
        daily_records.setdefault(str(row["sku_code"]), []).append(
            (str(row["sale_date"]), float(row["daily_qty"] or 0.0))
        )

    for sku_code, sku_df in data.monthly_df.groupby("sku_code", sort=False):
        sku_name = (sku_df["sku_name"].dropna().iloc[0] if sku_df["sku_name"].notna().any() else "") or ""
        sku_variant_key = (
            sku_df["sku_variant_key"].dropna().iloc[0] if sku_df["sku_variant_key"].notna().any() else ""
        ) or ""
        month_map = {
            str(row["year_month"]): float(row["monthly_qty"] or 0.0)
            for _, row in sku_df.iterrows()
            if pd.notna(row["year_month"])
        }
        sku_daily = daily_records.get(str(sku_code), [])
        current_month_actual = sum(
            qty for sale_date, qty in sku_daily if sale_date.startswith(data.latest_sale_month)
        )
        current_month_mean_forecast_qty = compute_current_month_mean_forecast_qty(
            sku_daily,
            data.latest_sale_month,
        )
        recent_7_day_qty = sum(qty for _, qty in sku_daily[-7:])
        recent_15_day_qty = sum(qty for _, qty in sku_daily[-15:])
        two_years_ago_total, previous_year_total, current_year_total = compute_actual_year_summaries(
            month_map,
            actual_history_months,
            data.latest_complete_month,
            data.latest_sale_month,
            current_month_actual,
        )

        total_future = 0.0
        month_values: dict[str, float] = {}
        growth_values: dict[str, float | None] = {}
        previous_qty_for_growth = float(month_map.get(data.latest_complete_month, 0.0))
        first_month_qty = build_default_month1_forecast(
            month_map=month_map,
            current_month_actual=current_month_actual,
            recent_7_day_qty=recent_7_day_qty,
            recent_15_day_qty=recent_15_day_qty,
            latest_sale_date=data.latest_sale_date,
            latest_sale_month=data.latest_sale_month,
            latest_complete_month=data.latest_complete_month,
        )
        month_values[all_future_months[0]] = first_month_qty
        growth_values[f"{all_future_months[0]}环比增幅(%)"] = format_pct_display(
            compute_mom_growth_pct(
                first_month_qty,
                previous_qty_for_growth,
            )
        )
        total_future += first_month_qty

        previous_month_qty = first_month_qty
        previous_target_month = all_future_months[0]
        for future_month in all_future_months[1:]:
            previous_hist_month = month_shift(previous_target_month, -12)
            current_hist_month = month_shift(future_month, -12)
            ratio = safe_ratio(
                month_map.get(current_hist_month, 0.0),
                month_map.get(previous_hist_month, 0.0),
                fallback=1.0,
            )
            qty = float(ceil_non_negative(max(previous_month_qty * ratio, 0.0)))
            month_values[future_month] = qty
            growth_values[f"{future_month}环比增幅(%)"] = format_pct_display(
                compute_mom_growth_pct(qty, previous_month_qty)
            )
            total_future += qty
            previous_month_qty = qty
            previous_target_month = future_month

        row: dict[str, object] = {
            "商品编码": str(sku_code),
            "商品名称": str(sku_name),
            "商品标签分类": extract_product_root(sku_variant_key),
            "最近12个月最高销量(不含本月)": compute_recent_12m_max_qty(month_map, data.latest_sale_month),
            "最近12个月平均销量(不含本月)": ceil_non_negative(
                compute_recent_12m_avg_qty(month_map, data.latest_sale_month)
            ),
            "7-15天销量参考": compute_short_term_reference(
                recent_7_day_qty, recent_15_day_qty, data.latest_sale_month
            ),
            "历史同期月份比例推算值": ceil_non_negative(
                compute_last_month_yoy_mom_reference(
                    month_map, data.latest_sale_month, data.latest_complete_month
                )
            ),
        }
        for actual_month in actual_history_months:
            row[actual_month] = round(float(month_map.get(actual_month, 0.0)), 4)
        if current_month_mean_forecast_column:
            row[current_month_mean_forecast_column] = current_month_mean_forecast_qty
        row[two_years_ago_total_column] = two_years_ago_total
        row[previous_year_total_column] = previous_year_total
        row[current_year_total_column] = current_year_total
        for future_month in all_future_months:
            row[future_month] = month_values[future_month]
            row[f"{future_month}环比增幅(%)"] = growth_values[f"{future_month}环比增幅(%)"]
        row["预测合计"] = int(total_future)
        rows.append(row)

    result = pd.DataFrame(rows)
    if not result.empty:
        result = append_product_root_summary_rows(
            result,
            actual_history_months
            + ([current_month_mean_forecast_column] if current_month_mean_forecast_column else [])
            + [two_years_ago_total_column, previous_year_total_column, current_year_total_column]
            + all_future_months
            + ["预测合计"],
            actual_history_months + all_future_months,
            all_future_months,
            "预测合计",
        )
    return result


def generate_standard_formula_forecast(
    data: LoadedSalesData, formula: str, months: int
) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    latest_month = data.latest_complete_month
    all_future_months = [month_shift(latest_month, step) for step in range(1, months + 1)]
    actual_history_months = build_month_sequence("2024-01", latest_month)
    current_month_mean_forecast_column = build_current_month_mean_forecast_column_name(
        data.latest_sale_month, latest_month
    )
    two_years_ago_total_column, previous_year_total_column, current_year_total_column = build_actual_summary_column_names(latest_month)
    daily_records: dict[str, list[tuple[str, float]]] = {}
    for _, row in data.daily_df.iterrows():
        daily_records.setdefault(str(row["sku_code"]), []).append(
            (str(row["sale_date"]), float(row["daily_qty"] or 0.0))
        )

    for sku_code, sku_df in data.monthly_df.groupby("sku_code", sort=False):
        sku_name = (sku_df["sku_name"].dropna().iloc[0] if sku_df["sku_name"].notna().any() else "") or ""
        sku_variant_key = (
            sku_df["sku_variant_key"].dropna().iloc[0] if sku_df["sku_variant_key"].notna().any() else ""
        ) or ""
        month_map = {str(row["year_month"]): float(row["monthly_qty"] or 0) for _, row in sku_df.iterrows() if pd.notna(row["year_month"])}
        sku_daily = daily_records.get(str(sku_code), [])
        current_month_actual = sum(
            qty for sale_date, qty in sku_daily if sale_date.startswith(data.latest_sale_month)
        )
        current_month_mean_forecast_qty = compute_current_month_mean_forecast_qty(
            sku_daily,
            data.latest_sale_month,
        )
        recent_7_day_qty = sum(qty for _, qty in sku_daily[-7:])
        recent_15_day_qty = sum(qty for _, qty in sku_daily[-15:])
        two_years_ago_total, previous_year_total, current_year_total = compute_actual_year_summaries(
            month_map,
            actual_history_months,
            latest_month,
            data.latest_sale_month,
            current_month_actual,
        )
        history = [month_map.get(month_shift(latest_month, -offset), 0.0) for offset in range(11, -1, -1)]

        total_future = 0.0
        month_values: dict[str, float] = {}
        growth_values: dict[str, float | None] = {}
        previous_qty_for_growth = float(month_map.get(latest_month, 0.0))
        for future_month in all_future_months:
            env = build_safe_env(history)
            qty = float(ceil_non_negative(eval_formula(formula, env)))
            history.append(qty)
            total_future += qty
            month_values[future_month] = qty
            growth_values[f"{future_month}环比增幅(%)"] = format_pct_display(
                compute_mom_growth_pct(
                    qty,
                    previous_qty_for_growth,
                )
            )
            previous_qty_for_growth = qty

        row: dict[str, object] = {
            "商品编码": str(sku_code),
            "商品名称": str(sku_name),
            "商品标签分类": extract_product_root(sku_variant_key),
            "最近12个月最高销量(不含本月)": compute_recent_12m_max_qty(month_map, latest_month),
            "最近12个月平均销量(不含本月)": ceil_non_negative(
                compute_recent_12m_avg_qty(month_map, latest_month)
            ),
            "7-15天销量参考": compute_short_term_reference(
                recent_7_day_qty, recent_15_day_qty, data.latest_sale_month
            ),
            "历史同期月份比例推算值": ceil_non_negative(
                compute_last_month_yoy_mom_reference(
                    month_map, data.latest_sale_month, latest_month
                )
            ),
        }
        for actual_month in actual_history_months:
            row[actual_month] = round(float(month_map.get(actual_month, 0.0)), 4)
        if current_month_mean_forecast_column:
            row[current_month_mean_forecast_column] = current_month_mean_forecast_qty
        row[two_years_ago_total_column] = two_years_ago_total
        row[previous_year_total_column] = previous_year_total
        row[current_year_total_column] = current_year_total
        for future_month in all_future_months:
            row[future_month] = month_values[future_month]
            row[f"{future_month}环比增幅(%)"] = growth_values[f"{future_month}环比增幅(%)"]
        row["预测合计"] = int(total_future)
        rows.append(row)

    result = pd.DataFrame(rows)
    if not result.empty:
        result = append_product_root_summary_rows(
            result,
            actual_history_months
            + ([current_month_mean_forecast_column] if current_month_mean_forecast_column else [])
            + [two_years_ago_total_column, previous_year_total_column, current_year_total_column]
            + all_future_months
            + ["预测合计"],
            actual_history_months + all_future_months,
            all_future_months,
            "预测合计",
        )
    return result


def generate_forecast(data: LoadedSalesData, formula: str, months: int) -> pd.DataFrame:
    if formula == BUILTIN_DEFAULT_FORMULA:
        return generate_builtin_default_forecast(data, months)
    return generate_standard_formula_forecast(data, formula, months)


def build_production_plan_df(
    forecast_df: pd.DataFrame,
    data: LoadedSalesData,
    months: int,
    target_turnover_days: int = 45,
) -> pd.DataFrame:
    if forecast_df.empty:
        return pd.DataFrame()

    plan_months = [month_shift(data.latest_sale_month, step) for step in range(months)]
    detail_rows = forecast_df[~forecast_df["商品编码"].isin(["汇总", "环比", "同比"])].copy()
    if detail_rows.empty:
        return pd.DataFrame()

    latest_sale_day = datetime.strptime(data.latest_sale_date, "%Y-%m-%d").date()
    latest_month_day_count = int(month_end(data.latest_sale_month)[-2:])
    remaining_days = max(latest_month_day_count - latest_sale_day.day, 0)

    inventory_map: dict[str, float | None] = {}
    in_transit_map: dict[str, float | None] = {}
    for sku_code, sku_df in data.monthly_df.groupby("sku_code", sort=False):
        latest_inventory = sku_df["latest_inventory"].dropna().iloc[0] if sku_df["latest_inventory"].notna().any() else None
        latest_in_transit = sku_df["latest_in_transit"].dropna().iloc[0] if sku_df["latest_in_transit"].notna().any() else None
        inventory_map[str(sku_code)] = float(latest_inventory) if latest_inventory is not None else None
        in_transit_map[str(sku_code)] = float(latest_in_transit) if latest_in_transit is not None else None

    plan_rows: list[dict[str, object]] = []

    def compute_covered_days(supply: float, periods: list[tuple[int, float, float]]) -> float:
        covered_days = 0.0
        remaining_supply = supply
        for period_days, period_demand, daily_usage in periods:
            if remaining_supply <= 0:
                break
            if period_demand <= 0 or daily_usage <= 0:
                covered_days += float(period_days)
                continue
            if remaining_supply >= period_demand:
                covered_days += float(period_days)
                remaining_supply -= period_demand
                continue
            covered_days += remaining_supply / daily_usage
            remaining_supply = 0.0
            break
        return covered_days

    def compute_required_qty_for_target_days(target_days: int, supply: float, periods: list[tuple[int, float, float]]) -> float:
        remaining_target_days = float(target_days)
        required_qty = 0.0
        for period_days, period_demand, daily_usage in periods:
            if remaining_target_days <= 0:
                break
            days_to_cover = min(float(period_days), remaining_target_days)
            if period_demand <= 0 or daily_usage <= 0:
                remaining_target_days -= days_to_cover
                continue
            required_qty += daily_usage * days_to_cover
            remaining_target_days -= days_to_cover
        return float(ceil_to_batch(max(required_qty - supply, 0.0), 50))

    for _, row in detail_rows.iterrows():
        sku_code = str(row["商品编码"])
        latest_inventory = float(inventory_map.get(sku_code) or 0.0)
        latest_in_transit = float(in_transit_map.get(sku_code) or 0.0)
        available_supply = latest_inventory + latest_in_transit
        coverage_periods: list[tuple[int, float, float]] = []
        plan_row: dict[str, object] = {
            "商品编码": sku_code,
            "商品名称": row["商品名称"],
            "最新库存数据": inventory_map.get(sku_code),
            "最新采购在途数据": in_transit_map.get(sku_code),
        }
        current_month_remaining_demand = 0.0
        for index, month in enumerate(plan_months):
            month_number = int(month.split("-")[1])
            month_qty = float(row.get(month, 0.0) or 0.0)
            month_day_count = int(month_end(month)[-2:])
            if index == 0:
                column_name = f"{month_number}月剩余需求"
                remaining_demand = float(
                    ceil_non_negative(month_qty * remaining_days / latest_month_day_count)
                )
                plan_row[column_name] = remaining_demand
                current_month_remaining_demand = remaining_demand
                current_daily_usage = round(month_qty / latest_month_day_count, 4) if latest_month_day_count else 0.0
                coverage_periods.append((remaining_days, remaining_demand, current_daily_usage))
                plan_row[f"{month_number}月销售差额"] = round(
                    ceil_non_negative(remaining_demand - available_supply),
                    0,
                )
            else:
                plan_row[f"{month_number}月每日用量"] = round(month_qty / month_day_count, 4) if month_day_count else 0.0
                column_name = f"{month_number}月需求"
                month_demand = month_qty
                plan_row[column_name] = month_demand
                coverage_periods.append((month_day_count, month_demand, float(plan_row[f"{month_number}月每日用量"])))
                if index == 1:
                    first_month_number = int(plan_months[0].split("-")[1])
                    combined_gap = float(
                        ceil_non_negative(current_month_remaining_demand + month_demand - available_supply)
                    )
                    plan_row[f"{first_month_number}月+{month_number}月销售差额"] = combined_gap
                    plan_row["采购在途是否满足需求"] = "是" if combined_gap == 0 else "否"
        total_five_month_demand = round(sum(period_demand for _, period_demand, _ in coverage_periods), 4)
        last_forecast_month = plan_months[-1]
        last_month_number = int(last_forecast_month.split("-")[1])
        combined_covered_days = compute_covered_days(available_supply, coverage_periods)
        for label, supply in [
            ("成品周转天数", latest_inventory),
            ("在途+库存合计可周转天数", available_supply),
        ]:
            if supply >= total_five_month_demand:
                plan_row[label] = (
                    f"{last_forecast_month}月底以后" if label == "成品周转天数" else f"{last_month_number}月底以后"
                )
                continue
            plan_row[label] = f"{int(round(compute_covered_days(supply, coverage_periods)))}天"
        plan_row["库存计划量"] = (
            compute_required_qty_for_target_days(target_turnover_days, available_supply, coverage_periods)
            if combined_covered_days < float(target_turnover_days)
            else 0.0
        )
        plan_rows.append(plan_row)
    result = pd.DataFrame(plan_rows)
    if result.empty:
        return result

    first_month_number = int(plan_months[0].split("-")[1])
    demand_columns = [
        f"{first_month_number}月剩余需求",
        *[f"{int(month.split('-')[1])}月需求" for month in plan_months[1:]],
    ]
    gap_columns = [f"{first_month_number}月销售差额"]
    if len(plan_months) > 1:
        second_month_number = int(plan_months[1].split("-")[1])
        gap_columns.append(f"{first_month_number}月+{second_month_number}月销售差额")
    daily_usage_columns = [f"{int(month.split('-')[1])}月每日用量" for month in plan_months[1:]]

    ordered_columns = [
        "商品编码",
        "商品名称",
        "最新库存数据",
        *[column for column in demand_columns if column in result.columns],
        *[column for column in gap_columns if column in result.columns],
        *[column for column in daily_usage_columns if column in result.columns],
        "最新采购在途数据",
        "采购在途是否满足需求",
        "成品周转天数",
        "在途+库存合计可周转天数",
        "库存计划量",
    ]
    return result.reindex(columns=[column for column in ordered_columns if column in result.columns])


def is_year_month_column(column: str) -> bool:
    parts = column.split("-")
    return (
        len(parts) == 2
        and len(parts[0]) == 4
        and len(parts[1]) == 2
        and parts[0].isdigit()
        and parts[1].isdigit()
    )


def reorder_forecast_columns(
    forecast_df: pd.DataFrame,
    latest_sale_month: str,
    latest_complete_month: str,
) -> pd.DataFrame:
    if forecast_df.empty:
        return forecast_df

    columns = list(forecast_df.columns)
    future_months = sorted(
        [column for column in columns if is_year_month_column(column) and f"{column}环比增幅(%)" in columns]
    )
    actual_months = sorted(
        [column for column in columns if is_year_month_column(column) and column not in future_months]
    )
    total_column = "预测合计" if "预测合计" in columns else next(
        (column for column in columns if column.startswith("未来") and column.endswith("个月总销量")),
        None,
    )
    growth_columns = [f"{month}环比增幅(%)" for month in future_months if f"{month}环比增幅(%)" in columns]
    current_month_mean_forecast_column = build_current_month_mean_forecast_column_name(
        latest_sale_month, latest_complete_month
    )
    metric_columns = [
        "最近12个月最高销量(不含本月)",
        "最近12个月平均销量(不含本月)",
        "7-15天销量参考",
        "历史同期月份比例推算值",
    ]
    year_order = [
        latest_complete_month[:4],
        str(int(latest_complete_month[:4]) - 1),
        str(int(latest_complete_month[:4]) - 2),
    ]
    two_years_ago_total_column, previous_year_total_column, current_year_total_column = build_actual_summary_column_names(
        latest_complete_month
    )
    summary_column_map = {
        year_order[0]: current_year_total_column,
        year_order[1]: previous_year_total_column,
        year_order[2]: two_years_ago_total_column,
    }

    ordered_columns: list[str] = ["商品编码", "商品名称"]
    ordered_columns.extend(future_months)
    if total_column:
        ordered_columns.append(total_column)
    ordered_columns.extend(growth_columns)
    ordered_columns.extend([column for column in metric_columns if column in columns])
    for year in year_order:
        summary_column = summary_column_map.get(year)
        if summary_column in columns:
            ordered_columns.append(summary_column)
        ordered_columns.extend([column for column in actual_months if column.startswith(year)])
        if (
            current_month_mean_forecast_column
            and current_month_mean_forecast_column.startswith(year)
            and current_month_mean_forecast_column in columns
        ):
            ordered_columns.append(current_month_mean_forecast_column)
    if "商品标签分类" in columns:
        ordered_columns.append("商品标签分类")

    remaining_columns = [column for column in columns if column not in ordered_columns]
    return forecast_df.reindex(columns=ordered_columns + remaining_columns)


def format_worksheet(ws: Worksheet) -> None:
    thin_side = Side(style="thin", color="000000")
    solid_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    default_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    text_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    max_lengths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value in (None, ""):
                continue
            display = str(cell.value)
            max_lengths[cell.column] = max(max_lengths.get(cell.column, 0), len(display))
            cell.border = solid_border
            cell.alignment = text_alignment if cell.column in (1, 2) else default_alignment
            if cell.row == 1:
                cell.font = header_font
                cell.fill = header_fill

    for column_index, max_length in max_lengths.items():
        width = min(max(max_length * 1.4, 10), 36)
        ws.column_dimensions[openpyxl.utils.get_column_letter(column_index)].width = width

    ws.sheet_view.showGridLines = False


def apply_forecast_highlights(ws: Worksheet) -> None:
    headers = [cell.value for cell in ws[1]]
    if not headers:
        return

    recent_max_column = "最近12个月最高销量(不含本月)"
    short_term_column = "7-15天销量参考"
    historical_ratio_column = "历史同期月份比例推算值"
    current_month_mean_forecast_index = next(
        (
            index + 1
            for index, header in enumerate(headers)
            if isinstance(header, str) and header.endswith("均值预测销量")
        ),
        None,
    )
    future_months = [
        header
        for header in headers
        if isinstance(header, str) and is_year_month_column(header) and f"{header}环比增幅(%)" in headers
    ]
    if recent_max_column not in headers:
        return

    recent_max_index = headers.index(recent_max_column) + 1
    growth_indexes = {
        f"{month}环比增幅(%)": headers.index(f"{month}环比增幅(%)") + 1
        for month in future_months
        if f"{month}环比增幅(%)" in headers
    }
    month_indexes = {month: headers.index(month) + 1 for month in future_months}
    highlight_fill = PatternFill(fill_type="solid", fgColor="FF0000")
    short_term_index = headers.index(short_term_column) + 1 if short_term_column in headers else None
    historical_ratio_index = headers.index(historical_ratio_column) + 1 if historical_ratio_column in headers else None

    for row in ws.iter_rows(min_row=2):
        sku_code = row[0].value
        if sku_code in ("汇总", "环比", "同比", None, ""):
            continue

        metric_cells: list[tuple[object, float]] = []
        for column_index in (
            current_month_mean_forecast_index,
            historical_ratio_index,
            short_term_index,
        ):
            if not column_index:
                continue
            metric_cell = row[column_index - 1]
            try:
                metric_value = float(metric_cell.value) if metric_cell.value not in (None, "") else None
            except (TypeError, ValueError):
                metric_value = None
            if metric_value is not None:
                metric_cells.append((metric_cell, metric_value))
        if metric_cells:
            max_metric_value = max(metric_value for _, metric_value in metric_cells)
            for metric_cell, metric_value in metric_cells:
                if metric_value == max_metric_value:
                    metric_cell.fill = highlight_fill

        recent_max_value = row[recent_max_index - 1].value
        try:
            recent_max_qty = float(recent_max_value) if recent_max_value not in (None, "") else None
        except (TypeError, ValueError):
            recent_max_qty = None

        for month, column_index in month_indexes.items():
            month_cell = row[column_index - 1]
            try:
                month_qty = float(month_cell.value) if month_cell.value not in (None, "") else None
            except (TypeError, ValueError):
                month_qty = None
            if recent_max_qty is not None and month_qty is not None and month_qty > recent_max_qty:
                month_cell.font = Font(color="FF0000", bold=month_cell.font.bold)

            growth_cell = row[growth_indexes[f"{month}环比增幅(%)"] - 1]
            growth_value = growth_cell.value
            if not isinstance(growth_value, str) or not growth_value.endswith("%"):
                continue
            try:
                growth_pct = int(growth_value[:-1])
            except ValueError:
                continue
            if growth_pct > 100 or growth_pct < -100:
                growth_cell.font = Font(color="FF0000", bold=growth_cell.font.bold)


def apply_forecast_header_fills(ws: Worksheet) -> None:
    light_gray_fill = PatternFill(fill_type="solid", fgColor="E7E6E6")
    light_yellow_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    light_blue_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    light_green_fill = PatternFill(fill_type="solid", fgColor="E2F0D9")
    headers = [cell.value for cell in ws[1]]

    for index, header in enumerate(headers, start=1):
        if not isinstance(header, str):
            continue
        if header.endswith("环比增幅(%)"):
            ws.cell(row=1, column=index).fill = light_gray_fill
            continue
        if header == "最近12个月最高销量(不含本月)":
            ws.cell(row=1, column=index).fill = light_yellow_fill
            continue
        if header == "最近12个月平均销量(不含本月)":
            ws.cell(row=1, column=index).fill = light_blue_fill
            continue
        if header.endswith("年销量汇总"):
            ws.cell(row=1, column=index).fill = light_gray_fill
            continue
        if header.endswith("年实际销量汇总"):
            ws.cell(row=1, column=index).fill = light_green_fill


def apply_forecast_summary_row_fills(ws: Worksheet) -> None:
    summary_fill = PatternFill(fill_type="solid", fgColor="E7E6E6")
    for row in ws.iter_rows(min_row=2):
        row_label = row[0].value
        if row_label not in ("汇总", "环比", "同比"):
            continue
        for cell in row:
            if cell.value in (None, ""):
                continue
            cell.fill = summary_fill


def write_excel(
    forecast_df: pd.DataFrame,
    production_plan_df: pd.DataFrame,
    latest_label: str,
    latest_sale_date: str,
    inventory_metrics_updated_at: str | None,
    latest_sale_month: str,
    latest_complete_month: str,
    formula_text: str,
    normalized_formula: str,
    months: int,
    top_k: int,
    displayed_top_k: int,
    target_turnover_days: int,
    product_category: str | None,
    sku_scope_label: str,
    job_id: str,
) -> Path:
    output_dir = Path(__file__).resolve().parent.parent / "outputs" / "final"
    output_dir.mkdir(parents=True, exist_ok=True)
    now = datetime.now().strftime("%Y%m%d-%H%M%S")
    file_name = f"formula-forecast-{now}-{job_id}.xlsx"
    file_path = output_dir / file_name

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "预测结果"
    if not forecast_df.empty:
        forecast_df = reorder_forecast_columns(
            forecast_df,
            latest_sale_month=latest_sale_month,
            latest_complete_month=latest_complete_month,
        )
        headers = list(forecast_df.columns)
        ws.append(headers)
        for row in forecast_df.itertuples(index=False, name=None):
            ws.append(list(row))
    else:
        ws["A1"] = "无预测结果"
    format_worksheet(ws)
    apply_forecast_header_fills(ws)
    apply_forecast_summary_row_fills(ws)
    apply_forecast_highlights(ws)

    plan = wb.create_sheet("生产计划")
    if not production_plan_df.empty:
        plan_headers = list(production_plan_df.columns)
        plan.append(plan_headers)
        for row in production_plan_df.itertuples(index=False, name=None):
            plan.append(list(row))
    else:
        plan["A1"] = "无生产计划结果"
    format_worksheet(plan)

    constants = wb.create_sheet("常量")
    constants_rows = [
        ["字段", "值"],
        ["销售数据截止日期", latest_sale_date],
        ["库存/在途字段最近写入时间", inventory_metrics_updated_at or "未更新"],
        ["库存计划量可周转天数", target_turnover_days],
    ]
    for row in constants_rows:
        constants.append(row)
    format_worksheet(constants)

    info = wb.create_sheet("公式说明")
    info_rows = [
        ["原始公式描述", formula_text],
        ["归一化公式", normalized_formula],
        ["预测月数", months],
        ["SKU范围来源", sku_scope_label],
        ["Top SKU 数量", displayed_top_k],
        ["品类过滤", product_category or "全部"],
        ["历史截止月份", latest_label],
    ]
    if normalized_formula == BUILTIN_DEFAULT_FORMULA:
        info_rows.append(
            [
                "变量说明",
                "builtin_default_5m=第1个月取历史比例推算值/本月累计销量折算值/短期销量折算值三者最大值, 后续月份按去年同期相邻月份销量比递推",
            ]
        )
    else:
        info_rows.append(
            [
                "变量说明",
                "m1=上个月销量, m2=上2个月销量, avg3/avg6/avg12=最近N个月平均销量, sum3/sum6/sum12=最近N个月总销量, trend=m1-m2, growth=(m1-m2)/m2",
            ]
        )
    for key, value in info_rows:
        info.append([key, value])
    format_worksheet(info)

    wb.save(file_path)
    cleanup_final_outputs(output_dir)
    return file_path


def main() -> None:
    args = parse_args()
    sku_filter_codes = load_sku_codes_from_excel(args.sku_list_excel) if args.sku_list_excel else None
    target_turnover_days = max(1, min(int(args.target_turnover_days), 365))
    data = load_top_history(
        args.db,
        args.top_k,
        args.top_basis,
        args.product_category,
        sku_filter_codes=sku_filter_codes,
    )
    forecast_df = generate_forecast(data, args.formula, args.months)
    production_plan_df = build_production_plan_df(
        forecast_df,
        data,
        args.months,
        target_turnover_days=target_turnover_days,
    )
    latest_label = data.latest_sale_month if args.formula == BUILTIN_DEFAULT_FORMULA else data.latest_complete_month
    sku_scope_label = "文件中的 SKU" if sku_filter_codes else f"Top {args.top_k} SKU"
    displayed_top_k = len(sku_filter_codes) if sku_filter_codes else args.top_k
    file_path = write_excel(
        forecast_df=forecast_df,
        production_plan_df=production_plan_df,
        latest_label=latest_label,
        latest_sale_date=data.latest_sale_date,
        inventory_metrics_updated_at=data.inventory_metrics_updated_at,
        latest_sale_month=data.latest_sale_month,
        latest_complete_month=data.latest_complete_month,
        formula_text=args.formula_text,
        normalized_formula=args.formula,
        months=args.months,
        top_k=args.top_k,
        displayed_top_k=displayed_top_k,
        target_turnover_days=target_turnover_days,
        product_category=args.product_category,
        sku_scope_label=sku_scope_label,
        job_id=args.job_id,
    )
    print(
        json.dumps(
            {
                "file_path": str(file_path),
                "file_name": file_path.name,
                "summary_sheet": "预测结果",
                "formula_sheet": "公式说明",
                "latest_month": latest_label,
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
