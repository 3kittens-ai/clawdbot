#!/usr/bin/env python3
from __future__ import annotations

import argparse
import ast
import json
import math
import sqlite3
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
import pandas as pd

FINAL_OUTPUT_KEEP_COUNT = 50


def cleanup_final_outputs(output_dir: Path, keep_count: int = FINAL_OUTPUT_KEEP_COUNT) -> None:
    files = [path for path in output_dir.iterdir() if path.is_file()]
    if len(files) <= keep_count:
        return
    files.sort(key=lambda path: path.stat().st_mtime, reverse=True)
    for old_file in files[keep_count:]:
        old_file.unlink(missing_ok=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate formula-based sales forecast Excel")
    parser.add_argument("--db", required=True)
    parser.add_argument("--formula", required=True)
    parser.add_argument("--formula-text", required=True)
    parser.add_argument("--months", type=int, required=True)
    parser.add_argument("--top-k", type=int, required=True)
    parser.add_argument("--job-id", required=True)
    return parser.parse_args()


def month_shift(year_month: str, offset: int) -> str:
    year, month = [int(part) for part in year_month.split("-")]
    total = (year * 12 + (month - 1)) + offset
    next_year = total // 12
    next_month = total % 12 + 1
    return f"{next_year:04d}-{next_month:02d}"


def month_end(year_month: str) -> str:
    next_month = month_shift(year_month, 1)
    next_month_start = datetime.strptime(f"{next_month}-01", "%Y-%m-%d").date()
    return (next_month_start - timedelta(days=1)).strftime("%Y-%m-%d")


def resolve_latest_complete_month(latest_sale_date: str) -> str:
    latest_sale_day = datetime.strptime(latest_sale_date, "%Y-%m-%d").date()
    latest_month = latest_sale_date[:7]
    if latest_sale_date == month_end(latest_month):
        return latest_month
    return month_shift(latest_month, -1)


def build_safe_env(history: Sequence[float]) -> dict[str, float]:
    latest = list(history)[-12:]
    padded = [0.0] * max(0, 12 - len(latest)) + latest

    def avg(count: int) -> float:
        bucket = padded[-count:]
        return float(sum(bucket) / count) if count else 0.0

    def total(count: int) -> float:
        return float(sum(padded[-count:])) if count else 0.0

    m1 = float(padded[-1])
    m2 = float(padded[-2])
    return {
        **{f"m{i}": float(padded[-i]) for i in range(1, 13)},
        "avg3": avg(3),
        "avg6": avg(6),
        "avg12": avg(12),
        "sum3": total(3),
        "sum6": total(6),
        "sum12": total(12),
        "trend": float(m1 - m2),
        "growth": float((m1 - m2) / m2) if m2 else 0.0,
    }


ALLOWED_FUNCS = {
    "min": min,
    "max": max,
    "round": round,
}


def eval_formula(expr: str, env: dict[str, float]) -> float:
    node = ast.parse(expr, mode="eval")

    def _eval(current: ast.AST) -> float:
        if isinstance(current, ast.Expression):
            return _eval(current.body)
        if isinstance(current, ast.Constant) and isinstance(current.value, (int, float)):
            return float(current.value)
        if isinstance(current, ast.Name):
            if current.id not in env:
                raise ValueError(f"unsupported variable: {current.id}")
            return float(env[current.id])
        if isinstance(current, ast.BinOp):
            left = _eval(current.left)
            right = _eval(current.right)
            if isinstance(current.op, ast.Add):
                return left + right
            if isinstance(current.op, ast.Sub):
                return left - right
            if isinstance(current.op, ast.Mult):
                return left * right
            if isinstance(current.op, ast.Div):
                return left / right if right else 0.0
            if isinstance(current.op, ast.Pow):
                return left**right
            raise ValueError("unsupported operator")
        if isinstance(current, ast.UnaryOp):
            operand = _eval(current.operand)
            if isinstance(current.op, ast.USub):
                return -operand
            if isinstance(current.op, ast.UAdd):
                return operand
            raise ValueError("unsupported unary operator")
        if isinstance(current, ast.Call) and isinstance(current.func, ast.Name):
            func = ALLOWED_FUNCS.get(current.func.id)
            if func is None:
                raise ValueError(f"unsupported function: {current.func.id}")
            args = [_eval(arg) for arg in current.args]
            return float(func(*args))
        raise ValueError("unsupported expression")

    value = _eval(node)
    if not math.isfinite(value):
        raise ValueError("formula result is not finite")
    return max(value, 0.0)


def load_top_history(db_path: str, top_k: int) -> tuple[pd.DataFrame, str]:
    conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True, timeout=5)
    conn.execute("PRAGMA query_only = ON")
    latest_sale_date = pd.read_sql_query(
        "SELECT MAX(sale_date) AS latest_sale_date FROM sales",
        conn,
    ).iloc[0]["latest_sale_date"]
    latest_month = resolve_latest_complete_month(str(latest_sale_date))
    history_start = month_shift(latest_month, -11)
    history_start_date = f"{history_start}-01"

    top_df = pd.read_sql_query(
        """
        WITH recent_sales AS (
          SELECT sale_date, barcode, sales_volume
          FROM sales INDEXED BY idx_sales_date
          WHERE sale_date >= ?
        ),
        top_skus AS (
          SELECT barcode AS sku_code, SUM(sales_volume) AS total_qty
          FROM recent_sales
          GROUP BY barcode
          ORDER BY total_qty DESC
          LIMIT ?
        )
        SELECT
          t.sku_code,
          d.sku_name,
          d.variant_key AS sku_variant_key,
          substr(r.sale_date, 1, 7) AS year_month,
          SUM(r.sales_volume) AS monthly_qty,
          t.total_qty
        FROM top_skus t
        LEFT JOIN dim_sku d ON d.sku_code = t.sku_code
        LEFT JOIN recent_sales r ON r.barcode = t.sku_code
        GROUP BY t.sku_code, d.sku_name, d.variant_key, substr(r.sale_date, 1, 7), t.total_qty
        ORDER BY t.total_qty DESC, t.sku_code, year_month
        """,
        conn,
        params=[history_start_date, top_k],
    )
    conn.close()
    return top_df, str(latest_month)


def extract_product_root(raw_variant_key: object) -> str:
    value = str(raw_variant_key or "").strip()
    if not value:
        return "未分类"
    parts = value.split("|")
    if len(parts) <= 4:
        return value
    root = "|".join(parts[:-4]).strip("|")
    return root or value


def append_product_root_summary_rows(
    ordered: pd.DataFrame,
    month_columns: list[str],
    total_column: str,
) -> pd.DataFrame:
    if ordered.empty:
        return ordered

    root_column = "商品标签分类"
    code_column = "商品编码"
    name_column = "商品名称"
    ordered = ordered.copy()
    ordered[root_column] = ordered[root_column].fillna("未分类").replace("", "未分类")

    root_totals = (
        ordered.groupby(root_column, as_index=False)[total_column].sum().sort_values(
            by=[total_column, root_column],
            ascending=[False, True],
        )
    )
    sections: list[pd.DataFrame] = []
    for product_root in root_totals[root_column]:
        group = ordered[ordered[root_column] == product_root].copy()
        group = group.sort_values(by=[total_column, code_column], ascending=[False, True])
        summary_row: dict[str, object] = {
            root_column: product_root,
            code_column: "汇总",
            name_column: f"{product_root} 汇总",
        }
        for month_column in month_columns:
            summary_row[month_column] = float(group[month_column].sum())
        summary_row[total_column] = float(group[total_column].sum())
        sections.append(pd.DataFrame([summary_row], columns=ordered.columns))
        sections.append(group)

    return pd.concat(sections, ignore_index=True)


def generate_forecast(df: pd.DataFrame, latest_month: str, formula: str, months: int) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    all_future_months = [month_shift(latest_month, step) for step in range(1, months + 1)]

    for sku_code, sku_df in df.groupby("sku_code", sort=False):
        sku_name = (sku_df["sku_name"].dropna().iloc[0] if sku_df["sku_name"].notna().any() else "") or ""
        sku_variant_key = (
            sku_df["sku_variant_key"].dropna().iloc[0] if sku_df["sku_variant_key"].notna().any() else ""
        ) or ""
        month_map = {str(row["year_month"]): float(row["monthly_qty"] or 0) for _, row in sku_df.iterrows() if pd.notna(row["year_month"])}
        history = [month_map.get(month_shift(latest_month, -offset), 0.0) for offset in range(11, -1, -1)]

        total_future = 0.0
        month_values: dict[str, float] = {}
        for future_month in all_future_months:
            env = build_safe_env(history)
            qty = eval_formula(formula, env)
            qty = round(qty, 4)
            history.append(qty)
            total_future += qty
            month_values[future_month] = qty

        row: dict[str, object] = {
            "商品编码": str(sku_code),
            "商品名称": str(sku_name),
            "商品标签分类": extract_product_root(sku_variant_key),
        }
        row.update(month_values)
        row[f"未来{months}个月总销量"] = round(total_future, 4)
        rows.append(row)

    result = pd.DataFrame(rows)
    if not result.empty:
        month_columns = [
            column
            for column in result.columns
            if column not in {"商品编码", "商品名称", "商品标签分类", f"未来{months}个月总销量"}
        ]
        result = append_product_root_summary_rows(result, month_columns, f"未来{months}个月总销量")
    return result


def write_excel(
    forecast_df: pd.DataFrame,
    latest_month: str,
    formula_text: str,
    normalized_formula: str,
    months: int,
    top_k: int,
    job_id: str,
) -> Path:
    output_dir = Path(__file__).resolve().parent.parent / "outputs" / "final"
    output_dir.mkdir(parents=True, exist_ok=True)
    now = datetime.now().strftime("%Y%m%d-%H%M%S")
    file_name = f"formula-forecast-{now}-{job_id}.xlsx"
    file_path = output_dir / file_name

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "预测结果"
    if not forecast_df.empty:
        headers = list(forecast_df.columns)
        ws.append(headers)
        for row in forecast_df.itertuples(index=False, name=None):
            ws.append(list(row))
    else:
        ws["A1"] = "无预测结果"

    info = wb.create_sheet("公式说明")
    info_rows = [
        ["原始公式描述", formula_text],
        ["归一化公式", normalized_formula],
        ["预测月数", months],
        ["Top SKU 数量", top_k],
        ["历史截止月份", latest_month],
        ["变量说明", "m1=上个月销量, m2=上2个月销量, avg3/avg6/avg12=最近N个月平均销量, sum3/sum6/sum12=最近N个月总销量, trend=m1-m2, growth=(m1-m2)/m2"],
    ]
    for key, value in info_rows:
        info.append([key, value])

    wb.save(file_path)
    cleanup_final_outputs(output_dir)
    return file_path


def main() -> None:
    args = parse_args()
    df, latest_month = load_top_history(args.db, args.top_k)
    forecast_df = generate_forecast(df, latest_month, args.formula, args.months)
    file_path = write_excel(
        forecast_df=forecast_df,
        latest_month=latest_month,
        formula_text=args.formula_text,
        normalized_formula=args.formula,
        months=args.months,
        top_k=args.top_k,
        job_id=args.job_id,
    )
    print(
        json.dumps(
            {
                "file_path": str(file_path),
                "file_name": file_path.name,
                "summary_sheet": "预测结果",
                "formula_sheet": "公式说明",
                "latest_month": latest_month,
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
