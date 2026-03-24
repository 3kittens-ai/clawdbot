#!/usr/bin/env python3
from __future__ import annotations

import argparse
import importlib
import json
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from forecasting.config import DB_PATH, OUTPUT_DIR  # noqa: E402

FINAL_OUTPUT_KEEP_COUNT = 50


def cleanup_final_outputs(output_dir: Path, keep_count: int = FINAL_OUTPUT_KEEP_COUNT) -> None:
    files = [path for path in output_dir.iterdir() if path.is_file()]
    if len(files) <= keep_count:
        return
    files.sort(key=lambda path: path.stat().st_mtime, reverse=True)
    for old_file in files[keep_count:]:
        old_file.unlink(missing_ok=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run forecasting jobs for Feishu workflows")
    parser.add_argument("--action", required=True, choices=["backtest", "train", "predict"])
    parser.add_argument("--job-id", required=True)
    parser.add_argument("--skip-run", action="store_true")
    return parser.parse_args()


def run_forecasting_action(action: str) -> None:
    module_name = {
        "backtest": "forecasting.evaluate",
        "train": "forecasting.train",
        "predict": "forecasting.predict",
    }[action]
    module = importlib.import_module(module_name)
    module.main()


def load_sku_names() -> pd.DataFrame:
    conn = sqlite3.connect(DB_PATH)
    try:
        df = pd.read_sql_query(
            "SELECT sku_code AS variant_key, sku_name, variant_key AS sku_variant_key FROM dim_sku",
            conn,
        )
        df["variant_key"] = df["variant_key"].astype(str)
        df["商品标签分类"] = df["sku_variant_key"].map(extract_product_root)
        return df
    finally:
        conn.close()


def extract_product_root(raw_variant_key: object) -> str:
    value = str(raw_variant_key or "").strip()
    if not value:
        return "未分类"
    parts = value.split("|")
    if len(parts) <= 4:
        return value
    root = "|".join(parts[:-4]).strip("|")
    return root or value


def append_product_root_summary_rows(
    ordered: pd.DataFrame,
    month_columns: list[str],
    total_column: str,
) -> pd.DataFrame:
    if ordered.empty:
        return ordered

    root_column = "商品标签分类"
    code_column = "商品编码"
    name_column = "商品名称"
    ordered = ordered.copy()
    ordered[root_column] = ordered[root_column].fillna("未分类").replace("", "未分类")

    root_totals = (
        ordered.groupby(root_column, as_index=False)[total_column].sum().sort_values(
            by=[total_column, root_column],
            ascending=[False, True],
        )
    )
    sections: list[pd.DataFrame] = []
    for product_root in root_totals[root_column]:
        group = ordered[ordered[root_column] == product_root].copy()
        group = group.sort_values(by=[total_column, code_column], ascending=[False, True])
        summary_row: dict[str, object] = {
            root_column: product_root,
            code_column: "汇总",
            name_column: f"{product_root} 汇总",
        }
        for month_column in month_columns:
            summary_row[month_column] = float(group[month_column].sum())
        summary_row[total_column] = float(group[total_column].sum())
        sections.append(pd.DataFrame([summary_row], columns=ordered.columns))
        sections.append(group)

    return pd.concat(sections, ignore_index=True)


def load_input_range() -> tuple[str, str]:
    conn = sqlite3.connect(DB_PATH)
    try:
        row = conn.execute(
            "SELECT MIN(substr(sale_date, 1, 7)) AS start_month, MAX(substr(sale_date, 1, 7)) AS end_month FROM sales"
        ).fetchone()
    finally:
        conn.close()
    return str(row[0]), str(row[1])


def create_output_dir() -> Path:
    output_dir = BASE_DIR / "outputs" / "final"
    output_dir.mkdir(parents=True, exist_ok=True)
    return output_dir


def create_timestamped_output_path(prefix: str, suffix: str, job_id: str) -> Path:
    now = datetime.now().strftime("%Y%m%d-%H%M%S")
    return create_output_dir() / f"{prefix}-{now}-{suffix}-{job_id}.xlsx"


def format_number(value: float) -> str:
    return f"{float(value):,.2f}".rstrip("0").rstrip(".")


def format_pct(value: float) -> str:
    return f"{float(value) * 100:.2f}%"


def load_backtest_outputs() -> dict[str, pd.DataFrame]:
    output_root = Path(OUTPUT_DIR)
    return {
        "eval": pd.read_csv(output_root / "backtest_eval.csv"),
        "global": pd.read_csv(output_root / "metrics_global.csv"),
        "horizon": pd.read_csv(output_root / "metrics_horizon.csv"),
        "family": pd.read_csv(output_root / "metrics_by_family.csv"),
        "category": pd.read_csv(output_root / "metrics_by_category.csv"),
        "channel": pd.read_csv(output_root / "metrics_by_channel.csv"),
        "province": pd.read_csv(output_root / "metrics_by_province.csv"),
        "feature_top": pd.read_csv(output_root / "feature_importance_top.csv"),
        "feature_tail": pd.read_csv(output_root / "feature_importance_tail.csv"),
    }


def calculate_top100_summary(eval_df: pd.DataFrame) -> dict[str, float]:
    sku_totals = eval_df.groupby("variant_key", as_index=False)["monthly_qty"].sum()
    top_keys = sku_totals.sort_values("monthly_qty", ascending=False).head(100)["variant_key"]
    top_df = eval_df[eval_df["variant_key"].isin(top_keys)].copy()
    actual_sum = float(top_df["monthly_qty"].sum())
    forecast_sum = float(top_df["forecast_qty"].sum())
    abs_err = float((top_df["monthly_qty"] - top_df["forecast_qty"]).abs().sum())
    return {
        "actual_sum": actual_sum,
        "forecast_sum": forecast_sum,
        "wape": abs_err / actual_sum if actual_sum else 0.0,
        "bias": (forecast_sum - actual_sum) / actual_sum if actual_sum else 0.0,
    }


def summarize_horizon_metrics(horizon_df: pd.DataFrame) -> str:
    lines = ["| Horizon | Actual | Forecast | WAPE | Bias |", "| --- | ---: | ---: | ---: | ---: |"]
    horizon_column = "forecast_horizon(month)"
    if "forecast_horizon_month" in horizon_df.columns:
        horizon_column = "forecast_horizon_month"
    for _, row in horizon_df.head(12).iterrows():
        horizon = int(row[horizon_column])
        lines.append(
            f"| {horizon} | {format_number(row['actual_sum'])} | "
            f"{format_number(row['forecast_sum'])} | {row['wape']:.4f} | "
            f"{format_pct(row['bias'])} |"
        )
    return "\n".join(lines)


def summarize_metric_table(df: pd.DataFrame, dimension: str, top_n: int = 5) -> str:
    lines = [
        f"| {dimension} | Actual | Forecast | WAPE | Bias |",
        "| --- | ---: | ---: | ---: | ---: |",
    ]
    for row in df.head(top_n).itertuples(index=False):
        lines.append(
            f"| {getattr(row, dimension)} | {format_number(getattr(row, 'actual_sum'))} | "
            f"{format_number(getattr(row, 'forecast_sum'))} | {getattr(row, 'wape'):.4f} | "
            f"{format_pct(getattr(row, 'bias'))} |"
        )
    return "\n".join(lines)


def summarize_features(df: pd.DataFrame, top_n: int = 10) -> str:
    lines = ["| Rank | Feature | Importance |", "| ---: | --- | ---: |"]
    for row in df.head(top_n).itertuples(index=False):
        lines.append(
            f"| {int(getattr(row, 'rank', 0) or 0)} | {getattr(row, 'feature_name')} | "
            f"{format_number(getattr(row, 'importance'))} |"
        )
    return "\n".join(lines)


def update_backtest_reports() -> dict[str, Path]:
    outputs = load_backtest_outputs()
    eval_df = outputs["eval"].copy()
    eval_df["variant_key"] = eval_df["variant_key"].astype(str)
    global_row = outputs["global"].iloc[0]
    top100 = calculate_top100_summary(eval_df)
    input_start, input_end = load_input_range()
    output_start = str(eval_df["year_month"].min())
    output_end = str(eval_df["year_month"].max())
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    top_features = outputs["feature_top"].head(10)["feature_name"].tolist()

    report_path = BASE_DIR / "MODEL_REPORT.md"
    executive_path = BASE_DIR / "MODEL_REPORT_EXECUTIVE_SUMMARY.md"
    appendix_path = BASE_DIR / "MODEL_REPORT_TECH_APPENDIX.md"

    report_body = f"""# 销售预测模型报告

## 1. 报告概览

- 生成时间：`{generated_at}`
- 数据输入范围：`{input_start}` 到 `data-base/sales_filtered.sqlite` 当前最新完整历史
- 本次回测输出范围：`{output_start}` 到 `{output_end}`
- 回测 SKU 数量：`{int(eval_df["variant_key"].nunique())}`
- 回测记录行数：`{int(len(eval_df))}`

## 2. 整体结果

- 实际销量：{format_number(global_row["actual_sum"])}
- 预测销量：{format_number(global_row["forecast_sum"])}
- WAPE：`{global_row["wape"]:.4f}`
- Bias：`{format_pct(global_row["bias"])}`
- RMSE：`{format_number(global_row["rmse"])}`

## 3. Top 100 SKU 汇总

- 实际销量：{format_number(top100["actual_sum"])}
- 预测销量：{format_number(top100["forecast_sum"])}
- WAPE：`{top100["wape"]:.4f}`
- Bias：`{format_pct(top100["bias"])}`

## 4. 分步表现

{summarize_horizon_metrics(outputs["horizon"])}

## 5. 主要特征

### 5.1 头部模型 Top 10

{summarize_features(outputs["feature_top"], 10)}

### 5.2 长尾模型 Top 10

{summarize_features(outputs["feature_tail"], 10)}

## 6. 重点切片

### 6.1 品类

{summarize_metric_table(outputs["category"], "product_category")}

### 6.2 渠道

{summarize_metric_table(outputs["channel"], "main_channel")}

### 6.3 省份

{summarize_metric_table(outputs["province"], "top_province")}

## 7. 结论

- 当前回测整体误差水平可直接用作经营预测参考，但仍不应把单个长尾 SKU 点预测当作刚性承诺。
- 当前影响力最高的特征主要集中在：{", ".join(top_features[:5])}。
- 建议持续跟踪 Bias 明显偏离的 family / channel / province 切片，并在专项校准后复跑完整回测。
"""

    executive_body = f"""# 销售预测模型报告（一页版）

## 1. 当前回测窗口

- 输入数据范围：`{input_start}` 到 `data-base/sales_filtered.sqlite` 当前最新完整历史
- 验证输出范围：`{output_start}` 到 `{output_end}`
- SKU 数量：`{int(eval_df["variant_key"].nunique())}`

## 2. 核心结果

- 整体实际销量：{format_number(global_row["actual_sum"])}
- 整体预测销量：{format_number(global_row["forecast_sum"])}
- 整体 WAPE：`{global_row["wape"]:.4f}`
- 整体 Bias：`{format_pct(global_row["bias"])}`

## 3. Top 100 SKU 结果

- 实际销量：{format_number(top100["actual_sum"])}
- 预测销量：{format_number(top100["forecast_sum"])}
- WAPE：`{top100["wape"]:.4f}`
- Bias：`{format_pct(top100["bias"])}`

## 4. 最重要的特征

{summarize_features(outputs["feature_top"], 6)}

## 5. 结论

- 这套模型已经可以用作月度经营预测工具。
- 当前最需要持续观察的不是单一总量，而是高销量 family、重点渠道和重点省份的方向性偏差。
- 详细分步指标、分切片指标和技术说明见 `MODEL_REPORT.md` 与 `MODEL_REPORT_TECH_APPENDIX.md`。
"""

    appendix_body = f"""# 销售预测模型技术附录

## 1. 本次回测产物

- `outputs/backtest_eval.csv`
- `outputs/metrics_global.csv`
- `outputs/metrics_horizon.csv`
- `outputs/metrics_by_category.csv`
- `outputs/metrics_by_family.csv`
- `outputs/metrics_by_channel.csv`
- `outputs/metrics_by_province.csv`
- `outputs/feature_importance_top.csv`
- `outputs/feature_importance_tail.csv`

## 2. 回测口径

- 输入历史范围：`{input_start}` 到 `data-base/sales_filtered.sqlite` 当前最新完整历史
- 验证输出范围：`{output_start}` 到 `{output_end}`
- 回测记录行数：`{int(len(eval_df))}`

## 3. 整体指标

| Metric | Value |
| --- | ---: |
| SKU Count | {int(global_row["sku_count"])} |
| Row Count | {int(global_row["row_count"])} |
| Actual Sum | {format_number(global_row["actual_sum"])} |
| Forecast Sum | {format_number(global_row["forecast_sum"])} |
| WAPE | {global_row["wape"]:.4f} |
| MAE | {format_number(global_row["mae"])} |
| MAPE | {global_row["mape"]:.4f} |
| RMSE | {format_number(global_row["rmse"])} |
| Bias | {format_pct(global_row["bias"])} |

## 4. 分步指标

{summarize_horizon_metrics(outputs["horizon"])}

## 5. Family Top 5

{summarize_metric_table(outputs["family"], "family_tags")}

## 6. Category Top 5

{summarize_metric_table(outputs["category"], "product_category")}

## 7. Channel Top 5

{summarize_metric_table(outputs["channel"], "main_channel")}

## 8. Province Top 5

{summarize_metric_table(outputs["province"], "top_province")}

## 9. 特征重要性

### 9.1 头部模型

{summarize_features(outputs["feature_top"], 10)}

### 9.2 长尾模型

{summarize_features(outputs["feature_tail"], 10)}
"""

    report_path.write_text(report_body, encoding="utf-8")
    executive_path.write_text(executive_body, encoding="utf-8")
    appendix_path.write_text(appendix_body, encoding="utf-8")

    return {
        "report": report_path,
        "executive": executive_path,
        "appendix": appendix_path,
    }


def create_backtest_excel(job_id: str) -> Path:
    backtest_eval_path = Path(OUTPUT_DIR) / "backtest_eval.csv"
    feature_path = Path(OUTPUT_DIR) / "feature_importance_top.csv"
    if not feature_path.exists():
        feature_path = Path(OUTPUT_DIR) / "feature_importance.csv"

    eval_df = pd.read_csv(backtest_eval_path)
    feature_df = pd.read_csv(feature_path)
    eval_df["variant_key"] = eval_df["variant_key"].astype(str)
    sku_df = load_sku_names()
    input_start, input_end = load_input_range()

    forecast_df = (
        eval_df.groupby(["variant_key", "year_month"], as_index=False)["forecast_qty"]
        .sum()
        .merge(sku_df, on="variant_key", how="left")
    )
    pivot_df = forecast_df.pivot_table(
        index=["variant_key", "sku_name", "商品标签分类"],
        columns="year_month",
        values="forecast_qty",
        aggfunc="sum",
        fill_value=0,
    ).reset_index()
    month_columns = [
        col for col in pivot_df.columns if col not in {"variant_key", "sku_name", "商品标签分类"}
    ]
    month_columns = sorted(month_columns)
    ordered = pivot_df[["variant_key", "sku_name", "商品标签分类", *month_columns]].copy()
    ordered["总销量"] = ordered[month_columns].sum(axis=1)
    ordered = ordered.rename(
        columns={"variant_key": "商品编码", "sku_name": "商品名称"},
    )
    ordered = append_product_root_summary_rows(ordered, month_columns, "总销量")

    wb = openpyxl.Workbook()
    summary_ws = wb.active
    summary_ws.title = "回测特征与范围"
    output_start = str(eval_df["year_month"].min())
    output_end = str(eval_df["year_month"].max())
    summary_rows = [
        ["输入数据开始月份", input_start],
        ["输入数据结束月份", input_end],
        ["回测输出开始月份", output_start],
        ["回测输出结束月份", output_end],
        ["回测 SKU 数量", int(eval_df["variant_key"].nunique())],
        ["回测记录行数", int(len(eval_df))],
        [],
        ["Top 10 重要特征", "", ""],
        ["排名", "特征名", "重要性"],
    ]
    for row in summary_rows:
        summary_ws.append(row)
    for row in feature_df.head(10).itertuples(index=False):
        rank = int(getattr(row, "rank", 0) or 0)
        feature_name = str(getattr(row, "feature_name"))
        importance = float(getattr(row, "importance"))
        summary_ws.append([rank, feature_name, importance])

    result_ws = wb.create_sheet("回测输出")
    result_ws.append(list(ordered.columns))
    for row in ordered.itertuples(index=False, name=None):
        result_ws.append(list(row))

    file_path = create_timestamped_output_path("forecasting-backtest", "backtest-export", job_id)
    wb.save(file_path)
    cleanup_final_outputs(file_path.parent)
    return file_path


def create_forecast_excel(job_id: str, action: str) -> Path:
    forecast_path = Path(OUTPUT_DIR) / "forecast_results.csv"
    forecast_df = pd.read_csv(forecast_path)
    forecast_df["variant_key"] = forecast_df["variant_key"].astype(str)
    sku_df = load_sku_names()
    merged = forecast_df.merge(sku_df, on="variant_key", how="left")
    pivot_df = merged.pivot_table(
        index=["variant_key", "sku_name", "商品标签分类"],
        columns="year_month",
        values="forecast_qty",
        aggfunc="sum",
        fill_value=0,
    ).reset_index()
    month_columns = [
        col for col in pivot_df.columns if col not in {"variant_key", "sku_name", "商品标签分类"}
    ]
    month_columns = sorted(month_columns)
    ordered = pivot_df[["variant_key", "sku_name", "商品标签分类", *month_columns]].copy()
    ordered["总销量"] = ordered[month_columns].sum(axis=1)
    ordered = ordered.rename(
        columns={"variant_key": "商品编码", "sku_name": "商品名称"},
    )
    ordered = append_product_root_summary_rows(ordered, month_columns, "总销量")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "预测输出"
    ws.append(list(ordered.columns))
    for row in ordered.itertuples(index=False, name=None):
        ws.append(list(row))

    # Keep predict/train outputs under outputs/final with the same business-facing
    # results-* naming convention expected by downstream users.
    file_path = create_timestamped_output_path("results", f"{action}-forecast-export", job_id)
    wb.save(file_path)
    cleanup_final_outputs(file_path.parent)
    return file_path


def main() -> None:
    args = parse_args()
    if not args.skip_run:
        run_forecasting_action(args.action)

    attachments: list[dict[str, str]] = []
    result: dict[str, object] = {"action": args.action}

    if args.action == "backtest":
        report_paths = update_backtest_reports()
        summary_path = report_paths["executive"]
        excel_path = create_backtest_excel(args.job_id)
        attachments.append({"path": str(summary_path), "file_name": summary_path.name})
        attachments.append({"path": str(excel_path), "file_name": excel_path.name})
        result["summary_markdown_path"] = str(summary_path)
        result["excel_path"] = str(excel_path)
    else:
        excel_path = create_forecast_excel(args.job_id, args.action)
        attachments.append({"path": str(excel_path), "file_name": excel_path.name})
        result["excel_path"] = str(excel_path)

    result["attachments"] = attachments
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
